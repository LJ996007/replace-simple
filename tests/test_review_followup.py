import unittest
from pathlib import Path
from tempfile import TemporaryDirectory
from unittest.mock import patch

from docx import Document
from openpyxl import Workbook, load_workbook
from openpyxl.cell.rich_text import CellRichText, TextBlock
from openpyxl.cell.text import InlineFont

import main
from replacement_rules import DELETE_MARKER
from string_replacer import replace_in_xlsx
from symbol_clause_extractor import SymbolClause, export_symbol_clauses_to_excel
from tender_info_extractor import extract_project_info_details
from word_table_exporter import ExportedTable, ExportedTableCell, export_word_tables_to_excel


class ReviewFollowupTests(unittest.TestCase):
    def test_excel_zero_matches_preserve_bytes_and_matched_rich_text_keeps_styles(self):
        with TemporaryDirectory() as directory:
            source = Path(directory) / "source.xlsx"
            output = Path(directory) / "output.xlsx"
            book = Workbook()
            rich = CellRichText(TextBlock(InlineFont(b=True), "prefix OL"),
                                TextBlock(InlineFont(i=True), "D tail"))
            book.active["A1"] = rich
            book.active["A2"] = CellRichText("plain ", TextBlock(InlineFont(b=True), "untouched"))
            book.active["A3"] = "OLD"
            book.active["A4"] = "=1+1"
            book.save(source)
            book.close()
            original = source.read_bytes()
            for destination in (source, output):
                self.assertEqual(replace_in_xlsx(str(source), [("absent", "new")], str(destination)), 0)
                self.assertEqual(destination.read_bytes(), original)
            self.assertEqual(replace_in_xlsx(str(source), [("OLD", "NEW"), ("NEW", "wrong")], str(output)), 2)
            result = load_workbook(output, rich_text=True)
            try:
                value = result.active["A1"].value
                self.assertEqual(str(value), "prefix NEW tail")
                self.assertTrue(value[0].font.b)
                self.assertTrue(value[1].font.i)
                self.assertEqual(result.active["A2"].value, CellRichText("plain ", TextBlock(InlineFont(b=True), "untouched")))
                self.assertEqual(result.active["A3"].value, "NEW")
                self.assertEqual(result.active["A4"].data_type, "f")
            finally:
                result.close()
            self.assertEqual(replace_in_xlsx(str(output), [("NEW", DELETE_MARKER)], str(output)), 2)
            result = load_workbook(output, rich_text=True)
            try:
                self.assertEqual(str(result.active["A1"].value), "prefix  tail")
                self.assertTrue(result.active["A1"].value[1].font.i)
            finally:
                result.close()

    def test_both_exports_preserve_destination_and_remove_partial_files_on_failure(self):
        table = ExportedTable([ExportedTableCell(1, 1, 1, 1, "text")], 1, 1)
        exporters = (
            lambda path: export_word_tables_to_excel("unused.docx", str(path), exported_tables=[table]),
            lambda path: export_symbol_clauses_to_excel("unused.docx", str(path), clauses=[SymbolClause("★", "text")]),
        )

        def interrupted_save(_book, path):
            Path(path).write_bytes(b"partial archive")
            raise OSError("disk full")

        for export in exporters:
            for existing in (False, True):
                for target, failure in (("openpyxl.workbook.workbook.Workbook.save", interrupted_save),
                                        ("file_io.os.replace", OSError("locked"))):
                    with self.subTest(existing=existing, target=target), TemporaryDirectory() as directory:
                        output = Path(directory) / "output.xlsx"
                        if existing:
                            output.write_bytes(b"original")
                        with patch(target, side_effect=failure) if isinstance(failure, Exception) else patch(target, failure):
                            with self.assertRaises(OSError):
                                export(output)
                        self.assertEqual(output.read_bytes() if output.exists() else None,
                                         b"original" if existing else None)
                        self.assertEqual(list(Path(directory).glob(".replace-simple-*")), [])

    def test_contact_split_requires_a_name_and_phone_and_avoids_conflicting_sources(self):
        cases = (
            (["采购人电话：010-12345678，010-87654321"], "", "010-12345678，010-87654321"),
            (["采购人联系人：张老师，李老师"], "张老师，李老师", ""),
            (["采购人电话：张老师，010-12345678"], "张老师", "010-12345678"),
            (["采购人联系人：张老师，13812345678"], "张老师", "13812345678"),
            (["采购人电话：张老师，010-12345678", "采购人电话：李老师，010-87654321"], "", ""),
        )
        with TemporaryDirectory() as directory:
            path = Path(directory) / "source.docx"
            for lines, contact, phone in cases:
                with self.subTest(lines=lines):
                    document = Document()
                    for line in lines:
                        document.add_paragraph(line)
                    document.save(path)
                    details = {item["field"]: item for item in extract_project_info_details(str(path))}
                    self.assertEqual(details["采购人联系人"]["value"], contact)
                    self.assertEqual(details["采购人电话"]["value"], phone)

    def test_failed_session_save_keeps_window_and_retry_closes_it(self):
        root = main.tk.Tk()
        root.withdraw()
        try:
            app = main.ReplaceSimpleApp(root, restore_session=False)
            app.rules_sheet.set_sheet_data([["OLD", "NEW"]])
            with patch.object(app, "_save_session", side_effect=OSError("disk full")), \
                    patch("main.messagebox.showerror") as error:
                app.on_close()
            self.assertTrue(root.winfo_exists())
            self.assertEqual(app.get_rules_from_table(), [("OLD", "NEW")])
            error.assert_called_once()
            app.on_close()
            self.assertEqual(root.tk.call("after", "info"), "")
        finally:
            try:
                root.destroy()
            except main.tk.TclError:
                pass

    def test_destroy_cancels_callbacks_without_cancelling_other_windows(self):
        root = main.tk.Tk()
        root.withdraw()
        try:
            app = main.ReplaceSimpleApp(root, restore_session=False)
            app.open_word_table_exporter()
            window = app.table_export_window
            survivor = root.after(60_000, lambda: None)
            child = main.ttk.Button(window.window)
            child_timer = child.after(60_000, lambda: None)
            window_timer = window.window.after(60_000, lambda: None)
            child.destroy()
            self.assertNotIn(child_timer, root.tk.call("after", "info"))
            window.close()
            pending = root.tk.call("after", "info")
            self.assertNotIn(window_timer, pending)
            self.assertIn(survivor, pending)
        finally:
            root.destroy()
        self.assertEqual(root.tk.call("after", "info"), "")

    def test_multi_row_toggle_calculates_visible_keys_once(self):
        from types import SimpleNamespace
        from word_export_window import WordTableExportWindow

        window = object.__new__(WordTableExportWindow)
        window.selected_scan_keys = set()
        window.scan_item_by_iid = {
            str(i): ("table", SimpleNamespace(file_path="sample.docx", table_index=i))
            for i in range(200)
        }
        with patch.object(window, "_visible_scan_keys", wraps=window._visible_scan_keys) as visible, \
                patch.object(window, "_apply_scan_filters"):
            window._toggle_scan_iids(list(window.scan_item_by_iid))
            self.assertEqual(len(window.selected_scan_keys), 200)
            visible.assert_called_once()
            window._toggle_scan_iids(list(window.scan_item_by_iid))
            self.assertFalse(window.selected_scan_keys)


if __name__ == "__main__":
    unittest.main()
