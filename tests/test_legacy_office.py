import os
import shutil
import tempfile
import unittest
from contextlib import contextmanager
from pathlib import Path
from unittest.mock import patch

import legacy_office
from string_replacer import batch_replace


class _FakePythonCom:
    def __init__(self):
        self.initialized = 0
        self.uninitialized = 0

    def CoInitialize(self):
        self.initialized += 1

    def CoUninitialize(self):
        self.uninitialized += 1


class _FakeApplication:
    def __init__(self, path=r"C:\Program Files\Microsoft Office", caption="Microsoft Office"):
        self.Path = path
        self.Caption = caption
        self.Name = caption
        self.quit_count = 0

    def Quit(self):
        self.quit_count += 1


class _FakeClient:
    def __init__(self, available):
        self.available = dict(available)
        self.attempts = []

    def DispatchEx(self, prog_id):
        self.attempts.append(prog_id)
        if prog_id not in self.available:
            raise RuntimeError("class not registered")
        return self.available[prog_id]


class _WordTextRoot:
    def __init__(self, text):
        self.text = text


class _FakeWordRange:
    def __init__(self, root, start=0, end=None):
        self.root = root
        self.Start = start
        self.End = len(root.text) if end is None else end

    @property
    def Text(self):
        return self.root.text[self.Start:self.End]

    @Text.setter
    def Text(self, value):
        self.root.text = self.root.text[:self.Start] + value + self.root.text[self.End:]
        self.End = self.Start + len(value)

    @property
    def Duplicate(self):
        return _FakeWordRange(self.root, self.Start, self.End)

    def SetRange(self, start, end):
        self.Start = start
        self.End = end


class LegacyOfficeTests(unittest.TestCase):
    def test_session_initializes_com_uses_fallback_and_quits(self):
        pythoncom = _FakePythonCom()
        application = _FakeApplication()
        client = _FakeClient({"KWPS.Application": application})

        with legacy_office.LegacyOfficeSession(pythoncom, client) as session:
            self.assertIs(session.application("word"), application)
            self.assertEqual(session.backend_name("word"), "WPS Writer")

        self.assertEqual(client.attempts[:2], ["Word.Application", "KWPS.Application"])
        self.assertEqual(application.quit_count, 1)
        self.assertEqual(pythoncom.initialized, 1)
        self.assertEqual(pythoncom.uninitialized, 1)

    def test_session_reports_missing_application(self):
        pythoncom = _FakePythonCom()
        client = _FakeClient({})

        with legacy_office.LegacyOfficeSession(pythoncom, client) as session:
            with self.assertRaises(legacy_office.OfficeApplicationUnavailableError):
                session.application("excel")

    def test_office_progid_registered_by_wps_is_identified_and_released_on_switch(self):
        pythoncom = _FakePythonCom()
        word = _FakeApplication(r"C:\Kingsoft\WPS Office\office6", "WPS Writer")
        excel = _FakeApplication(r"C:\Kingsoft\WPS Office\office6", "WPS Spreadsheets")
        client = _FakeClient({"Word.Application": word, "Excel.Application": excel})

        with legacy_office.LegacyOfficeSession(pythoncom, client) as session:
            session.application("word")
            self.assertEqual(session.backend_name("word"), "WPS Writer")
            session.application("excel")
            self.assertEqual(word.quit_count, 1)
            self.assertIsNone(session.backend_name("word"))
            self.assertEqual(session.backend_name("excel"), "WPS Spreadsheets")

        self.assertEqual(excel.quit_count, 1)

    def test_atomic_copy_replaces_output_only_after_success(self):
        with tempfile.TemporaryDirectory() as directory:
            source = Path(directory) / "source.doc"
            output = Path(directory) / "output.doc"
            source.write_bytes(b"original")
            output.write_bytes(b"previous")

            with legacy_office._atomic_editable_copy(str(source), str(output)) as temporary:
                Path(temporary).write_bytes(b"replaced")

            self.assertEqual(source.read_bytes(), b"original")
            self.assertEqual(output.read_bytes(), b"replaced")

    def test_atomic_copy_preserves_existing_output_on_failure(self):
        with tempfile.TemporaryDirectory() as directory:
            source = Path(directory) / "source.xls"
            output = Path(directory) / "output.xls"
            source.write_bytes(b"original")
            output.write_bytes(b"previous")

            with self.assertRaisesRegex(RuntimeError, "stop"):
                with legacy_office._atomic_editable_copy(str(source), str(output)) as temporary:
                    Path(temporary).write_bytes(b"partial")
                    raise RuntimeError("stop")

            self.assertEqual(source.read_bytes(), b"original")
            self.assertEqual(output.read_bytes(), b"previous")

    def test_word_range_replacement_is_non_cascading_and_supports_long_text(self):
        root = _WordTextRoot("AB")
        word_range = _FakeWordRange(root)
        replacement = "长" * 400
        prepared = legacy_office._prepared_rules([("A", replacement), ("B", "C"), ("长", "X")])

        count = legacy_office._replace_in_word_range(word_range, prepared)

        self.assertEqual(count, 2)
        self.assertEqual(root.text, replacement + "C")

    def test_batch_replace_routes_doc_to_legacy_backend(self):
        calls = []

        class FakeSession:
            def __enter__(self):
                calls.append("enter")
                return self

            def close(self):
                calls.append("close")

        def fake_replace(file_path, rules, output_path, session):
            calls.append((Path(file_path).suffix, tuple(rules), Path(output_path).suffix))
            shutil.copy2(file_path, output_path)
            return 3

        with tempfile.TemporaryDirectory() as directory:
            source = Path(directory) / "采购人.doc"
            output_dir = Path(directory) / "output"
            output_dir.mkdir()
            source.write_bytes(b"legacy")
            with (
                patch.object(legacy_office, "LegacyOfficeSession", FakeSession),
                patch.object(legacy_office, "replace_in_doc", fake_replace),
            ):
                results, error = batch_replace(
                    [str(source)],
                    [("采购人", "建设单位")],
                    output_dir=str(output_dir),
                )

        self.assertIsNone(error)
        self.assertEqual(results, {"采购人.doc": 3})
        self.assertEqual(calls[0], "enter")
        self.assertEqual(calls[-1], "close")
        self.assertEqual(calls[1][0], ".doc")

    def test_main_supported_extensions_include_all_legacy_formats(self):
        from main import SUPPORTED_EXTENSIONS, WORD_EXTENSIONS

        self.assertTrue({".doc", ".xls", ".ppt"}.issubset(SUPPORTED_EXTENSIONS))
        self.assertEqual(WORD_EXTENSIONS, (".doc", ".docx"))

    def test_doc_table_scan_uses_converted_docx_and_reports_original_path(self):
        from docx import Document
        from word_table_exporter import batch_scan_word_tables

        class FakeSession:
            def __enter__(self):
                return self

            def close(self):
                pass

        with tempfile.TemporaryDirectory() as directory:
            source = Path(directory) / "source.doc"
            converted = Path(directory) / "converted.docx"
            source.write_bytes(b"legacy")
            document = Document()
            table = document.add_table(rows=1, cols=1)
            table.cell(0, 0).text = "采购人"
            document.save(converted)

            @contextmanager
            def fake_source(_path, _session):
                yield str(converted)

            with (
                patch.object(legacy_office, "LegacyOfficeSession", FakeSession),
                patch.object(legacy_office, "temporary_docx_source", fake_source),
            ):
                items, skipped, error = batch_scan_word_tables([str(source)])

        self.assertIsNone(error)
        self.assertEqual(skipped, {})
        self.assertEqual(len(items), 1)
        self.assertEqual(items[0].file_path, str(source))
        self.assertEqual(items[0].filename, "source.doc")


if __name__ == "__main__":
    unittest.main()
