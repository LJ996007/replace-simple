import datetime
import sys
import time
import tkinter as tk
import tkinter.font as tkfont
import unittest
import zipfile
from pathlib import Path
from tempfile import TemporaryDirectory
from tkinter import ttk

import xlwt
from docx import Document
from docx.oxml import parse_xml
from docx.oxml.ns import qn
from openpyxl import Workbook, load_workbook
from pptx import Presentation

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

import main as simple_main
from main import (
    BackgroundTaskRunner,
    CappedScrollbarModel,
    DEFAULT_PRESETS,
    OUTPUT_DIR_HINT,
    ReplaceSimpleApp,
    append_presets_to_rule_rows,
    collect_supported_files,
    elide_middle,
    normalize_rule_rows,
    normalize_presets,
    presets_from_settings,
    remove_matching_file_paths,
)
from string_replacer import (
    apply_rules_to_filename,
    batch_replace,
    get_output_path,
    load_replacement_rules,
    replace_in_docx,
    replace_in_pptx,
)
from tender_info_extractor import extract_project_info, extract_project_info_rules
from word_table_exporter import (
    batch_export_word_tables,
    export_word_tables_to_excel,
    get_table_export_output_path,
    scan_word_tables,
)


def create_hidden_root():
    last_error = None
    for _attempt in range(3):
        try:
            root = tk.Tk()
            root.withdraw()
            return root
        except tk.TclError as exc:
            last_error = exc
            time.sleep(0.2)
    raise last_error


class SimpleReplacementTests(unittest.TestCase):
    def test_default_presets_and_saved_empty_list(self):
        self.assertEqual(len(DEFAULT_PRESETS), 12)
        self.assertIn("[项目名称]", DEFAULT_PRESETS)
        self.assertIn("[采购人地址]", DEFAULT_PRESETS)
        self.assertEqual(presets_from_settings({}), list(DEFAULT_PRESETS))
        self.assertEqual(presets_from_settings({"presets": []}), [])

    def test_normalize_presets_trims_and_dedupes(self):
        self.assertEqual(
            normalize_presets([" [项目名称] ", "", "[项目名称]", None, "[项目编号]"]),
            ["[项目名称]", "[项目编号]"],
        )

    def test_append_presets_to_rule_rows_skips_existing_old_text(self):
        rows, added = append_presets_to_rule_rows(
            [["[项目名称]", "已有项目"], ["", ""]],
            ["[项目名称]", "[项目编号]", "[项目编号]"],
        )
        self.assertEqual(added, 1)
        self.assertEqual(rows, [["[项目名称]", "已有项目"], ["[项目编号]", ""]])

    def test_background_task_runner_returns_events_on_tk_thread(self):
        root = create_hidden_root()
        try:
            runner = BackgroundTaskRunner(root)
            progress = []
            completed = []
            errors = []

            submitted = runner.submit(
                lambda publish: (publish("读取中"), "完成")[1],
                completed.append,
                errors.append,
                progress.append,
            )

            deadline = time.monotonic() + 1
            while not completed and time.monotonic() < deadline:
                root.update()
                time.sleep(0.01)

            self.assertTrue(submitted)
            self.assertFalse(runner.active)
            self.assertEqual(progress, ["读取中"])
            self.assertEqual(completed, ["完成"])
            self.assertEqual(errors, [])
        finally:
            root.destroy()

    def test_collect_supported_files_skips_temporary_and_unsupported_files(self):
        with TemporaryDirectory() as tmp_dir:
            root = Path(tmp_dir)
            nested = root / "nested"
            nested.mkdir()
            (root / "normal.docx").touch()
            (nested / "normal.xlsx").touch()
            (nested / "~$temporary.docx").touch()
            (nested / "notes.txt").touch()
            progress = []

            files = collect_supported_files(
                str(root),
                (".docx", ".xlsx"),
                progress.append,
            )

        self.assertEqual({Path(path).name for path in files}, {"normal.docx", "normal.xlsx"})
        self.assertTrue(progress)

    def test_load_replacement_rules_reads_two_columns_from_first_row(self):
        with TemporaryDirectory() as tmp_dir:
            rules_path = Path(tmp_dir) / "rules.xlsx"
            wb = Workbook()
            ws = wb.active
            ws["A1"] = "采购人"
            ws["B1"] = "建设单位"
            ws["A2"] = "供应商"
            ws["B2"] = "投标人"
            wb.save(rules_path)

            rules = load_replacement_rules(str(rules_path))

        self.assertEqual(rules, [("采购人", "建设单位"), ("供应商", "投标人")])

    def test_load_replacement_rules_supports_xlsm_files(self):
        with TemporaryDirectory() as tmp_dir:
            rules_path = Path(tmp_dir) / "rules.xlsm"
            wb = Workbook()
            ws = wb.active
            ws["A1"] = "采购人"
            ws["B1"] = "建设单位"
            wb.save(rules_path)

            rules = load_replacement_rules(str(rules_path))

        self.assertEqual(rules, [("采购人", "建设单位")])

    def test_load_replacement_rules_supports_xls_files(self):
        with TemporaryDirectory() as tmp_dir:
            rules_path = Path(tmp_dir) / "rules.xls"
            wb = xlwt.Workbook()
            ws = wb.add_sheet("Sheet1")
            ws.write(0, 0, "采购人")
            ws.write(0, 1, "建设单位")
            ws.write(1, 0, "供应商")
            ws.write(1, 1, "投标人")
            wb.save(str(rules_path))

            rules = load_replacement_rules(str(rules_path))

        self.assertEqual(rules, [("采购人", "建设单位"), ("供应商", "投标人")])

    def test_normalize_rule_rows_skips_empty_old_text_and_dedupes_exact_pairs(self):
        rules = normalize_rule_rows([
            (" 采购人 ", " 建设单位 "),
            ("", "忽略"),
            ("供应商", ""),
            ("采购人", "建设单位"),
            ("采购人", "采购单位"),
        ])

        self.assertEqual(rules, [
            ("采购人", "建设单位"),
            ("供应商", ""),
            ("采购人", "采购单位"),
        ])

    def test_app_starts_with_compact_fonts_and_blank_editable_rows(self):
        root = create_hidden_root()
        try:
            app = ReplaceSimpleApp(root, restore_session=False)
            root.update_idletasks()
            data = app.rules_sheet.get_sheet_data()

            self.assertLessEqual(app.body_font.cget("size"), 10)
            self.assertLessEqual(app.title_font.cget("size"), 13)
            self.assertGreaterEqual(len(data), 3)
            self.assertEqual(list(data[0]), ["", ""])
        finally:
            root.destroy()

    def test_elide_middle_keeps_head_and_tail_within_width(self):
        root = create_hidden_root()
        try:
            font = tkfont.Font(root, family="Microsoft YaHei UI", size=10)
            text = "C:/Users/LJ/WPSDrive/1154674488/WPS云盘/前期资料-公开招标"
            fitted = elide_middle(text, font, 180)
            self.assertIn("...", fitted)
            self.assertTrue(fitted.startswith("C:/"))
            self.assertTrue(fitted.endswith("招标"))
            self.assertLessEqual(font.measure(fitted), 180)
            self.assertEqual(elide_middle("abc", font, 1000), "abc")
        finally:
            root.destroy()

    def test_long_output_dir_does_not_push_start_button_away(self):
        root = create_hidden_root()
        try:
            app = ReplaceSimpleApp(root, restore_session=False)
            root.update_idletasks()
            long_dir = (
                "C:/Users/LJ/WPSDrive/1154674488/WPS云盘/"
                "#北航2026/BUAAZB20260105-北京航空航天大学软件学院高性能服务器/"
                "前期资料-公开招标"
            )
            app._apply_output_dir(long_dir)
            app.status_var.set("已选择输出目录")
            app._refresh_output_and_status_layout()
            root.update_idletasks()

            shown = app.output_label.cget("text")
            self.assertEqual(app.output_dir, long_dir)
            self.assertLessEqual(len(shown), len(long_dir))
            self.assertNotIn(long_dir, app.status_var.get())
            self.assertEqual(app.status_var.get(), "已选择输出目录")
            self.assertLessEqual(app.output_label.winfo_reqheight(), 70)
            self.assertGreater(app.start_button.winfo_reqwidth(), 50)
            wraplength = int(float(app.output_label.cget("wraplength") or 0))
            self.assertGreaterEqual(wraplength, 200)
            self.assertLessEqual(
                app.body_font.measure(shown),
                wraplength * 2 + 8,
            )
            self.assertNotEqual(shown, OUTPUT_DIR_HINT)
        finally:
            root.destroy()

    def test_long_status_and_progress_filename_do_not_cover_start_button(self):
        root = create_hidden_root()
        try:
            app = ReplaceSimpleApp(root, restore_session=False)
            root.update_idletasks()
            long_name = "BUAAZB20260105-北京航空航天大学软件学院高性能服务器公开招标文件" * 3 + ".docx"
            app.status_var.set(f"正在处理 (1/8)：{long_name}")
            app._refresh_output_and_status_layout()
            root.update_idletasks()

            shown = app.status_label.cget("text")
            self.assertLess(len(shown), len(app.status_var.get()))
            self.assertIn("...", shown)
            self.assertGreater(app.start_button.winfo_reqwidth(), 50)
            self.assertLessEqual(app.status_label.winfo_reqheight(), 40)
        finally:
            root.destroy()

    def test_table_exporter_long_path_and_status_stay_on_one_line(self):
        root = create_hidden_root()
        app = None
        try:
            app = ReplaceSimpleApp(root, restore_session=False)
            app.open_word_table_exporter()
            root.update_idletasks()
            exporter = app.table_export_window
            long_dir = (
                "C:/Users/LJ/WPSDrive/1154674488/WPS云盘/"
                "#北航2026/BUAAZB20260105-北京航空航天大学软件学院高性能服务器/"
                "前期资料-公开招标"
            )
            exporter.output_dir = long_dir
            exporter._output_path.set_text(long_dir, foreground="green")
            exporter.status_var.set(
                "正在导出 (1/2)：BUAAZB20260105-北京航空航天大学软件学院高性能服务器招标文件.docx"
            )
            exporter._refresh_constrained_texts()
            root.update_idletasks()

            self.assertLess(len(exporter.output_label.cget("text")), len(long_dir))
            self.assertIn("...", exporter.output_label.cget("text"))
            self.assertLessEqual(exporter.output_label.winfo_reqheight(), 40)
            self.assertLess(len(exporter.status_label.cget("text")), len(exporter.status_var.get()))
            self.assertGreater(exporter.export_button.winfo_reqwidth(), 50)
            exporter.close()
        finally:
            if app is not None and app.table_export_window is not None:
                app.table_export_window.close()
            root.destroy()

    def test_app_uses_flat_sections_without_label_frames(self):
        root = create_hidden_root()
        try:
            ReplaceSimpleApp(root, restore_session=False)

            def walk(widget):
                children = widget.winfo_children()
                for child in children:
                    yield child
                    yield from walk(child)

            label_frames = [child for child in walk(root) if isinstance(child, ttk.LabelFrame)]

            self.assertEqual(label_frames, [])
        finally:
            root.destroy()

    def test_app_does_not_enable_sv_ttk_theme_on_startup(self):
        calls = []

        class ThemeProbe:
            def set_theme(self, theme):
                calls.append(theme)

        had_sv_ttk = hasattr(simple_main, "sv_ttk")
        previous_sv_ttk = getattr(simple_main, "sv_ttk", None)
        simple_main.sv_ttk = ThemeProbe()

        root = create_hidden_root()
        try:
            ReplaceSimpleApp(root, restore_session=False)
            self.assertEqual(calls, [])
        finally:
            root.destroy()
            if had_sv_ttk:
                simple_main.sv_ttk = previous_sv_ttk
            else:
                delattr(simple_main, "sv_ttk")

    def test_flat_scrollbar_style_has_small_arrow_elements(self):
        root = create_hidden_root()
        try:
            ReplaceSimpleApp(root, restore_session=False)
            style = ttk.Style(root)
            layout_text = str(style.layout("Flat.Vertical.TScrollbar"))
            self.assertIn("thumb", layout_text)
            self.assertIn("uparrow", layout_text)
            self.assertIn("downarrow", layout_text)
            self.assertEqual(style.lookup("Flat.Vertical.TScrollbar", "troughcolor"), "#EDF3F8")
            self.assertEqual(style.lookup("Flat.Vertical.TScrollbar", "background"), "#BFD3E7")
            self.assertEqual(style.lookup("Flat.Vertical.TScrollbar", "arrowcolor"), "#5D7891")
        finally:
            root.destroy()

    def test_capped_scrollbar_limits_full_span_and_maps_drag(self):
        displayed = []
        commands = []
        model = CappedScrollbarModel(
            lambda *args: commands.append(args),
            lambda first, last: displayed.append((first, last)),
        )
        model.set(0, 1)
        self.assertEqual(displayed[-1], (0.14, 0.86))

        model.set(0.25, 0.75)
        self.assertEqual(displayed[-1], (0.25, 0.75))
        model.dispatch("moveto", 0.5)
        self.assertEqual(commands[-1], ("moveto", 0.5))

    def test_vertical_mousewheel_scrolls_in_both_directions(self):
        root = create_hidden_root()
        try:
            app = ReplaceSimpleApp(root, restore_session=False)

            class ScrollTarget:
                def __init__(self):
                    self.calls = []

                def yview_scroll(self, amount, unit):
                    self.calls.append((amount, unit))

            class WheelEvent:
                def __init__(self, delta):
                    self.delta = delta

            target = ScrollTarget()
            handler = app._bind_vertical_mousewheel(target)
            self.assertEqual(handler(WheelEvent(-120)), "break")
            self.assertEqual(handler(WheelEvent(120)), "break")
            self.assertEqual(target.calls, [(1, "units"), (-1, "units")])
        finally:
            root.destroy()

    def test_app_has_word_table_export_entry_and_opens_window(self):
        root = create_hidden_root()
        app = None
        try:
            app = ReplaceSimpleApp(root, restore_session=False)
            root.update_idletasks()

            self.assertEqual(app.table_export_button.cget("text"), "提取信息")
            self.assertEqual(app.version_info_button.cget("text"), "版本")
            self.assertEqual(root.title(), simple_main.format_version_title())

            app.version_info_button.invoke()
            root.update_idletasks()
            version_windows = [
                child
                for child in root.winfo_children()
                if isinstance(child, tk.Toplevel) and child.title() == "版本信息"
            ]
            self.assertEqual(len(version_windows), 1)
            version_windows[0].destroy()

            app.open_word_table_exporter()
            root.update_idletasks()

            self.assertIsNotNone(app.table_export_window)
            self.assertTrue(app.table_export_window.exists())
            self.assertEqual(app.table_export_window.window.title(), "提取信息")
            self.assertTrue(hasattr(app.table_export_window, "scan_tree"))
            self.assertEqual(app.table_export_window.scan_button.cget("text"), "扫描表格")
            self.assertGreaterEqual(app.table_export_window.WINDOW_WIDTH, 1200)
            self.assertGreaterEqual(app.table_export_window.WINDOW_HEIGHT, 880)
            self.assertGreaterEqual(int(app.table_export_window.scan_tree.cget("height")), 12)
            self.assertGreaterEqual(int(app.table_export_window.detail_text.cget("height")), 6)
            self.assertTrue(hasattr(app.table_export_window, "detail_scrollbar"))
            self.assertEqual(
                app.table_export_window.scan_tree.cget("columns"),
                ("selected", "file", "section", "table", "size"),
            )
            self.assertEqual(app.table_export_window.scan_tree.heading("section", "text"), "所在章节")
            self.assertNotIn("hint", app.table_export_window.scan_tree.cget("columns"))
            self.assertEqual(
                len(app.table_export_window._scan_column_separators),
                len(app.table_export_window.scan_tree.cget("columns")) - 1,
            )
            self.assertTrue(all(
                separator.cget("background") == simple_main.HEADER_GRID_COLOR
                for separator in app.table_export_window._scan_column_separators
            ))
            self.assertTrue(hasattr(app.table_export_window, "_scan_header_bottom_line"))
            self.assertEqual(
                app.table_export_window._scan_header_bottom_line.cget("background"),
                simple_main.HEADER_GRID_COLOR,
            )
            header_height = app.table_export_window._scan_tree_header_height()
            self.assertGreaterEqual(header_height, 20)

            exporter = app.table_export_window
            fitted_w, fitted_h, min_w, min_h = exporter._fitted_window_size()
            self.assertGreaterEqual(fitted_w, min_w)
            self.assertGreaterEqual(fitted_h, min_h)
            self.assertLessEqual(fitted_w, exporter.WINDOW_WIDTH)
            self.assertLessEqual(fitted_h, exporter.WINDOW_HEIGHT)
            exporter._sync_scan_tree_x_scrollbar(0.0, 1.0)
            exporter.window.update_idletasks()
            self.assertFalse(bool(exporter.scan_x_scrollbar.grid_info()))
            exporter._sync_scan_tree_x_scrollbar(0.0, 0.7)
            exporter.window.update_idletasks()
            self.assertTrue(bool(exporter.scan_x_scrollbar.grid_info()))
            exporter._sync_scan_tree_x_scrollbar(0.0, 1.0)
            exporter.window.update_idletasks()
            self.assertFalse(bool(exporter.scan_x_scrollbar.grid_info()))

            app.table_export_window.close()
            root.update_idletasks()
            self.assertIsNone(app.table_export_window)
        finally:
            if app is not None and app.table_export_window is not None:
                app.table_export_window.close()
            root.destroy()

    def test_app_has_tender_file_import_button(self):
        root = create_hidden_root()
        try:
            ReplaceSimpleApp(root, restore_session=False)

            def walk(widget):
                children = widget.winfo_children()
                for child in children:
                    yield child
                    yield from walk(child)

            button_texts = [
                child.cget("text")
                for child in walk(root)
                if isinstance(child, ttk.Button)
            ]

            self.assertIn("导入招标文件", button_texts)
            self.assertIn("预设 ▼", button_texts)
            self.assertNotIn("导出 Excel", button_texts)
        finally:
            root.destroy()

    def test_preset_popup_adds_multiple_selected_presets(self):
        root = create_hidden_root()
        try:
            app = ReplaceSimpleApp(root, restore_session=False)
            app.clear_rules()
            app._preset_vars = [
                ("[项目名称]", tk.BooleanVar(value=True)),
                ("[项目编号]", tk.BooleanVar(value=True)),
                ("[标的名称]", tk.BooleanVar(value=False)),
            ]
            app.add_selected_presets()
            self.assertEqual(
                app.get_rules_from_table(),
                [("[项目名称]", ""), ("[项目编号]", "")],
            )
        finally:
            root.destroy()

    def test_preset_manager_shows_edit_controls_without_clipping(self):
        root = create_hidden_root()
        try:
            app = ReplaceSimpleApp(root, restore_session=False)
            app.open_preset_manager()
            root.update_idletasks()
            dialogs = [child for child in root.winfo_children() if isinstance(child, tk.Toplevel)]
            self.assertEqual(len(dialogs), 1)
            dialog = dialogs[0]

            def walk(widget):
                for child in widget.winfo_children():
                    yield child
                    yield from walk(child)

            widgets = list(walk(dialog))
            button_texts = [
                child.cget("text") for child in widgets if isinstance(child, ttk.Button)
            ]
            self.assertIn("新增", button_texts)
            self.assertIn("保存修改", button_texts)
            self.assertIn("删除选中", button_texts)
            self.assertIn("保存并关闭", button_texts)
            self.assertTrue(any(isinstance(child, ttk.Entry) for child in widgets))
            self.assertEqual(dialog.minsize(), (520, 540))
            dialog.destroy()
        finally:
            root.destroy()

    def test_tender_import_reuses_remembered_word_file_without_dialog(self):
        with TemporaryDirectory() as tmp_dir:
            source_path = Path(tmp_dir) / "tender.docx"
            doc = Document()
            doc.add_paragraph("项目名称：一次选择项目")
            doc.save(source_path)

            root = create_hidden_root()
            previous_dialog = simple_main.filedialog.askopenfilename
            try:
                app = ReplaceSimpleApp(root, restore_session=False)
                app._remember_word_files([str(source_path)])

                def fail_dialog(*_args, **_kwargs):
                    raise AssertionError("不应重复弹出文件选择窗口")

                simple_main.filedialog.askopenfilename = fail_dialog
                app.import_rules_from_tender_file()

                deadline = time.monotonic() + 2
                while not app.get_rules_from_table() and time.monotonic() < deadline:
                    root.update()
                    time.sleep(0.01)

                self.assertIn(("[项目名称]", "一次选择项目"), app.get_rules_from_table())
            finally:
                simple_main.filedialog.askopenfilename = previous_dialog
                root.destroy()

    def test_word_table_exporter_reuses_remembered_word_file(self):
        with TemporaryDirectory() as tmp_dir:
            source_path = Path(tmp_dir) / "tender.docx"
            doc = Document()
            doc.add_paragraph("项目名称：表格复用项目")
            doc.save(source_path)

            root = create_hidden_root()
            app = None
            try:
                app = ReplaceSimpleApp(root, restore_session=False)
                app._remember_word_files([str(source_path)])

                app.open_word_table_exporter()
                root.update_idletasks()

                self.assertIsNotNone(app.table_export_window)
                self.assertEqual(app.table_export_window.file_paths, [str(source_path)])
            finally:
                if app is not None and app.table_export_window is not None:
                    app.table_export_window.close()
                root.destroy()

    def test_make_blank_rows_returns_independent_rows(self):
        from main import make_blank_rows

        rows = make_blank_rows(4)
        self.assertEqual(rows, [["", ""], ["", ""], ["", ""], ["", ""]])
        # 各行必须是独立对象：改第 0 行不应连带改写其它行
        rows[0][0] = "X"
        self.assertEqual([r[0] for r in rows], ["X", "", "", ""])
        self.assertEqual(make_blank_rows(0), [])

    def test_remove_matching_file_paths_removes_selected_rules_workbook(self):
        with TemporaryDirectory() as tmp_dir:
            tmp_path = Path(tmp_dir)
            target = tmp_path / "rules.xlsx"
            target.touch()
            other = tmp_path / "document.xlsx"
            other.touch()

            remaining, removed = remove_matching_file_paths(
                [str(other), str(target)],
                str(target),
            )

        self.assertEqual(remaining, [str(other)])
        self.assertEqual(removed, [str(target)])

    def test_remove_matching_file_paths_handles_equivalent_relative_paths(self):
        with TemporaryDirectory() as tmp_dir:
            tmp_path = Path(tmp_dir)
            target = tmp_path / "rules.xlsx"
            target.touch()
            other = tmp_path / "document.xlsx"
            other.touch()
            relative_target = target.parent / "." / target.name

            remaining, removed = remove_matching_file_paths(
                [str(other), str(relative_target)],
                str(target),
            )

        self.assertEqual(remaining, [str(other)])
        self.assertEqual(removed, [str(relative_target)])

    def test_get_output_path_replaces_filename_without_suffix(self):
        output_path = get_output_path(
            r"C:\work\采购人.xlsx",
            [("采购人", "建设单位")],
            None,
        )

        self.assertEqual(Path(output_path).name, "建设单位.xlsx")

    def test_get_output_path_overwrites_source_when_filename_would_not_change(self):
        output_path = get_output_path(
            r"C:\work\sample.xlsx",
            [("采购人", "建设单位")],
            None,
        )

        self.assertEqual(Path(output_path).name, "sample.xlsx")
        self.assertEqual(Path(output_path), Path(r"C:\work\sample.xlsx"))

    def test_get_output_path_overwrites_source_when_selected_output_dir_is_source_dir(self):
        output_path = get_output_path(
            r"C:\work\sample.xlsx",
            [("采购人", "建设单位")],
            r"C:\work",
        )

        self.assertEqual(Path(output_path).name, "sample.xlsx")
        self.assertEqual(Path(output_path), Path(r"C:\work\sample.xlsx"))

    def test_get_output_path_avoids_existing_non_source_output_file(self):
        with TemporaryDirectory() as tmp_dir:
            root = Path(tmp_dir)
            source_dir = root / "source"
            output_dir = root / "output"
            source_dir.mkdir()
            output_dir.mkdir()
            source_path = source_dir / "sample.xlsx"
            source_path.touch()
            existing_output = output_dir / "sample.xlsx"
            existing_output.touch()

            output_path = get_output_path(
                str(source_path),
                [("采购人", "建设单位")],
                output_dir=str(output_dir),
            )

        self.assertEqual(Path(output_path).name, "sample_1.xlsx")

    def test_filename_replacement_preserves_extension(self):
        filename = apply_rules_to_filename(
            "投标文件.docx",
            [("doc", "文档"), ("投标", "响应")],
        )

        self.assertEqual(filename, "响应文件.docx")

    def test_rules_are_applied_longest_old_text_first(self):
        filename = apply_rules_to_filename(
            "采购人名单.xlsx",
            [("采购", "购买"), ("采购人", "建设单位")],
        )

        self.assertEqual(filename, "建设单位名单.xlsx")

    def test_filename_replacement_does_not_repeat_when_new_text_contains_old_text(self):
        filename = apply_rules_to_filename(
            "北京航空航天大学安全保卫部学院路校区中控室大屏更换采购.xlsx",
            [("北京航空航天大学安全保卫部学院路校区中控室大屏更换", "北京航空航天大学安全保卫部学院路校区中控室大屏更换采购")],
        )

        self.assertEqual(filename, "北京航空航天大学安全保卫部学院路校区中控室大屏更换采购.xlsx")

    def test_filename_replacement_converts_windows_invalid_characters(self):
        filename = apply_rules_to_filename(
            "[项目编号]评标报告.docx",
            [("[项目编号]", '2641STC73032/01:*?"<>|\\')],
        )

        self.assertEqual(filename, "2641STC73032／01：＊？＂＜＞｜＼评标报告.docx")

    def test_batch_replace_keeps_slash_in_content_but_sanitizes_output_filename(self):
        with TemporaryDirectory() as tmp_dir:
            root = Path(tmp_dir)
            source_dir = root / "source"
            output_dir = root / "output"
            source_dir.mkdir()
            source_path = source_dir / "[项目编号]评标报告.docx"
            document = Document()
            document.add_paragraph("项目编号：[项目编号]")
            document.save(source_path)

            results, error = batch_replace(
                [str(source_path)],
                [("[项目编号]", "2641STC73032/01")],
                output_dir=str(output_dir),
            )

            output_path = output_dir / "2641STC73032／01评标报告.docx"
            output_document = Document(output_path)

            self.assertIsNone(error)
            self.assertEqual(results, {"[项目编号]评标报告.docx": 1})
            self.assertTrue(output_path.exists())
            self.assertFalse((output_dir / "2641STC73032").exists())
            self.assertEqual(output_document.paragraphs[0].text, "项目编号：2641STC73032/01")

    def test_batch_replace_writes_replaced_xlsx_in_original_dir(self):
        with TemporaryDirectory() as tmp_dir:
            source_path = Path(tmp_dir) / "采购人.xlsx"
            wb = Workbook()
            ws = wb.active
            ws["A1"] = "采购人名单"
            wb.save(source_path)

            results, error = batch_replace(
                [str(source_path)],
                [("采购人", "建设单位")],
                output_dir=None,
            )

            output_path = Path(tmp_dir) / "建设单位.xlsx"
            output_exists = output_path.exists()
            output_wb = load_workbook(output_path)
            try:
                value = output_wb.active["A1"].value
            finally:
                output_wb.close()

        self.assertIsNone(error)
        self.assertEqual(results, {"采购人.xlsx": 1})
        self.assertTrue(output_exists)
        self.assertEqual(value, "建设单位名单")

    def test_batch_replace_overwrites_source_when_filename_does_not_change(self):
        with TemporaryDirectory() as tmp_dir:
            source_path = Path(tmp_dir) / "sample.xlsx"
            wb = Workbook()
            ws = wb.active
            ws["A1"] = "采购人名单"
            wb.save(source_path)

            results, error = batch_replace([str(source_path)], [("采购人", "建设单位")], output_dir=None)

            self.assertFalse((Path(tmp_dir) / "sample_已替换.xlsx").exists())
            source_wb = load_workbook(source_path)
            try:
                source_value = source_wb.active["A1"].value
            finally:
                source_wb.close()

        self.assertIsNone(error)
        self.assertEqual(results, {"sample.xlsx": 1})
        self.assertEqual(source_value, "建设单位名单")

    def test_batch_replace_uses_distinct_outputs_for_same_filename(self):
        with TemporaryDirectory() as tmp_dir:
            root = Path(tmp_dir)
            first_dir = root / "first"
            second_dir = root / "second"
            output_dir = root / "output"
            first_dir.mkdir()
            second_dir.mkdir()
            output_dir.mkdir()
            source_paths = [first_dir / "sample.xlsx", second_dir / "sample.xlsx"]

            for source_path in source_paths:
                workbook = Workbook()
                workbook.active["A1"] = "采购人名单"
                workbook.save(source_path)
                workbook.close()

            results, error = batch_replace(
                [str(path) for path in source_paths],
                [("采购人", "建设单位")],
                output_dir=str(output_dir),
            )

            output_paths = [output_dir / "sample.xlsx", output_dir / "sample_1.xlsx"]
            output_values = []
            for output_path in output_paths:
                workbook = load_workbook(output_path)
                try:
                    output_values.append(workbook.active["A1"].value)
                finally:
                    workbook.close()

        self.assertIsNone(error)
        self.assertEqual(len(results), 2)
        self.assertEqual(output_values, ["建设单位名单", "建设单位名单"])

    def test_docx_cross_run_replacement_preserves_unmatched_run_formatting(self):
        with TemporaryDirectory() as tmp_dir:
            source_path = Path(tmp_dir) / "source.docx"
            output_path = Path(tmp_dir) / "output.docx"

            doc = Document()
            paragraph = doc.add_paragraph()
            first = paragraph.add_run("采购")
            first.bold = True
            second = paragraph.add_run("人")
            second.italic = True
            tail = paragraph.add_run("名单")
            tail.underline = True
            doc.save(source_path)

            count = replace_in_docx(str(source_path), [("采购人", "建设单位")], str(output_path))

            output_doc = Document(output_path)
            runs = output_doc.paragraphs[0].runs

        self.assertEqual(count, 1)
        self.assertEqual([run.text for run in runs], ["建设单位", "", "名单"])
        self.assertTrue(runs[0].bold)
        self.assertTrue(runs[2].underline)

    def test_docx_does_not_repeat_when_new_text_contains_old_text(self):
        with TemporaryDirectory() as tmp_dir:
            source_path = Path(tmp_dir) / "source.docx"
            output_path = Path(tmp_dir) / "output.docx"

            doc = Document()
            doc.add_paragraph("北京航空航天大学安全保卫部学院路校区中控室大屏更换采购")
            doc.save(source_path)

            count = replace_in_docx(
                str(source_path),
                [
                    (
                        "北京航空航天大学安全保卫部学院路校区中控室大屏更换",
                        "北京航空航天大学安全保卫部学院路校区中控室大屏更换采购",
                    )
                ],
                str(output_path),
            )

            output_doc = Document(output_path)
            text = output_doc.paragraphs[0].text

        self.assertEqual(count, 0)
        self.assertEqual(text, "北京航空航天大学安全保卫部学院路校区中控室大屏更换采购")

    def test_extract_project_info_rules_from_tender_docx(self):
        with TemporaryDirectory() as tmp_dir:
            source_path = Path(tmp_dir) / "tender.docx"

            doc = Document()
            doc.add_paragraph("第一章 投标邀请")
            doc.add_paragraph("项目名称：智慧校园设备采购项目")
            doc.add_paragraph("项目编号：ABC-2026-001")
            doc.add_paragraph("四、提交投标文件截止时间、开标时间和地点")
            doc.add_paragraph("投标截止时间、开标时间：2026 年 7 月 20 日 09:30（北京时间）。")
            doc.add_paragraph("地点：北京市海淀区1号会议室。")
            doc.add_paragraph("五、公告期限")
            table = doc.add_table(rows=3, cols=5)
            table.cell(0, 0).text = "包号"
            table.cell(0, 1).text = "标的名称"
            table.cell(0, 2).text = "数量"
            table.cell(1, 0).text = "01"
            table.cell(1, 1).text = "中控室大屏及显示控制系统"
            table.cell(1, 2).text = "1套"
            table.cell(2, 0).text = "02"
            table.cell(2, 1).text = "配套控制设备"
            table.cell(2, 2).text = "2套"
            doc.add_paragraph("1.采购人信息")
            doc.add_paragraph("名 称：示例大学")
            doc.add_paragraph("联系方式：张老师，010-12345678")
            doc.add_paragraph("地址：北京市海淀区学院路1号")
            doc.add_paragraph("示例招标代理机构")
            doc.add_paragraph("2026年7月1日")
            doc.add_paragraph("第二章 投标人须知")
            doc.save(source_path)

            info = extract_project_info(str(source_path))
            rules = extract_project_info_rules(str(source_path))

        self.assertEqual(info["项目名称"], "智慧校园设备采购项目")
        self.assertEqual(info["标的名称"], "中控室大屏及显示控制系统；配套控制设备")
        self.assertEqual(info["采购人名称"], "示例大学")
        self.assertEqual(info["采购人联系人"], "张老师")
        self.assertEqual(info["采购人电话"], "010-12345678")
        self.assertEqual(info["采购人地址"], "北京市海淀区学院路1号")
        self.assertEqual(info["开标时间"], "2026 年 7 月 20 日 09:30")
        self.assertEqual(info["开标日期"], "2026年7月20日")
        self.assertEqual(info["招标公告日期"], "2026年7月1日")
        self.assertEqual(
            rules,
            [
                ("[项目名称]", "智慧校园设备采购项目"),
                ("[项目编号]", "ABC-2026-001"),
                ("[标的名称]", "中控室大屏及显示控制系统；配套控制设备"),
                ("[采购人名称]", "示例大学"),
                ("[采购人联系人]", "张老师"),
                ("[采购人电话]", "010-12345678"),
                ("[开标时间]", "2026 年 7 月 20 日 09:30"),
                ("[开标日期]", "2026年7月20日"),
                ("[招标公告日期]", "2026年7月1日"),
                ("[开标地点]", "北京市海淀区1号会议室"),
                ("[报名人数]", ""),
                ("[采购人地址]", "北京市海淀区学院路1号"),
            ],
        )

    def test_extract_opening_date_tolerates_word_character_spacing(self):
        with TemporaryDirectory() as tmp_dir:
            source_path = Path(tmp_dir) / "spaced-opening-time.docx"

            doc = Document()
            doc.add_paragraph("开 标 时 间：2 0 2 6 年 0 7 月 2 0 日 13点00分（北 京 时 间）")
            doc.save(source_path)

            info = extract_project_info(str(source_path))
            rules = dict(extract_project_info_rules(str(source_path)))

        self.assertEqual(info["开标时间"], "2 0 2 6 年 0 7 月 2 0 日 13点00分")
        self.assertEqual(info["开标日期"], "2026年7月20日")
        self.assertEqual(rules["[开标日期]"], "2026年7月20日")

    def test_extract_project_info_handles_procurement_contact_context(self):
        with TemporaryDirectory() as tmp_dir:
            source_path = Path(tmp_dir) / "notice.docx"

            doc = Document()
            doc.add_paragraph("1.采购人信息")
            doc.add_paragraph("名 称：示例采购单位")
            doc.add_paragraph("联系人：采购人联系人")
            doc.add_paragraph("电 话：010-11112222")
            doc.add_paragraph("地 址：采购人地址")
            doc.add_paragraph("2.采购代理机构信息")
            doc.add_paragraph("名称：示例代理机构")
            doc.add_paragraph("联系人：代理联系人")
            doc.add_paragraph("电话：010-99998888")
            doc.save(source_path)

            info = extract_project_info(str(source_path))

        self.assertEqual(info["采购人名称"], "示例采购单位")
        self.assertEqual(info["采购人联系人"], "采购人联系人")
        self.assertEqual(info["采购人电话"], "010-11112222")
        self.assertEqual(info["采购人地址"], "采购人地址")

    def test_extract_project_info_handles_table_context_and_parenthesized_labels(self):
        with TemporaryDirectory() as tmp_dir:
            source_path = Path(tmp_dir) / "table-context.docx"

            doc = Document()
            table = doc.add_table(rows=6, cols=4)
            table.cell(0, 0).text = "采购人信息"
            table.cell(1, 0).text = "名称"
            table.cell(1, 1).text = "表格采购单位"
            table.cell(2, 0).text = "项目名称"
            table.cell(2, 1).text = "设备采购"
            table.cell(2, 2).text = "项目编号"
            table.cell(2, 3).text = "XYZ-2026"
            table.cell(3, 0).text = "联系人"
            table.cell(3, 1).text = "李老师"
            table.cell(4, 0).text = "联系电话"
            table.cell(4, 1).text = "010-87654321"
            table.cell(5, 0).text = "地址"
            table.cell(5, 1).text = "北京市朝阳区1号"
            doc.save(source_path)

            info = extract_project_info(str(source_path))

        self.assertEqual(info["采购人名称"], "表格采购单位")
        self.assertEqual(info["项目名称"], "设备采购")
        self.assertEqual(info["项目编号"], "XYZ-2026")
        self.assertEqual(info["采购人联系人"], "李老师")
        self.assertEqual(info["采购人电话"], "010-87654321")
        self.assertEqual(info["采购人地址"], "北京市朝阳区1号")

    def test_extract_project_info_handles_parenthesized_paragraph_labels(self):
        with TemporaryDirectory() as tmp_dir:
            source_path = Path(tmp_dir) / "parenthesized-paragraph.docx"

            doc = Document()
            doc.add_paragraph("（一）采购人名称（单位）：示例单位")
            doc.add_paragraph("2.开标日期（北京时间）：2026-07-20")
            doc.save(source_path)

            info = extract_project_info(str(source_path))

        self.assertEqual(info["采购人名称"], "示例单位")
        self.assertEqual(info["开标日期"], "2026-07-20")

    def test_extract_project_info_does_not_cross_into_next_field_value(self):
        with TemporaryDirectory() as tmp_dir:
            source_path = Path(tmp_dir) / "missing-value.docx"

            doc = Document()
            table = doc.add_table(rows=1, cols=3)
            table.cell(0, 0).text = "项目名称"
            table.cell(0, 1).text = "项目编号"
            table.cell(0, 2).text = "ABC-001"
            doc.save(source_path)

            info = extract_project_info(str(source_path))

        self.assertNotIn("项目名称", info)
        self.assertEqual(info["项目编号"], "ABC-001")

    def test_pptx_cross_run_replacement_preserves_unmatched_run_formatting(self):
        with TemporaryDirectory() as tmp_dir:
            source_path = Path(tmp_dir) / "source.pptx"
            output_path = Path(tmp_dir) / "output.pptx"

            presentation = Presentation()
            slide = presentation.slides.add_slide(presentation.slide_layouts[6])
            textbox = slide.shapes.add_textbox(0, 0, 3000000, 1000000)
            paragraph = textbox.text_frame.paragraphs[0]
            first = paragraph.add_run()
            first.text = "供"
            first.font.bold = True
            second = paragraph.add_run()
            second.text = "应商"
            second.font.italic = True
            tail = paragraph.add_run()
            tail.text = "名单"
            tail.font.underline = True
            presentation.save(source_path)

            count = replace_in_pptx(str(source_path), [("供应商", "投标人")], str(output_path))

            output_presentation = Presentation(output_path)
            output_shape = output_presentation.slides[0].shapes[0]
            runs = output_shape.text_frame.paragraphs[0].runs

        self.assertEqual(count, 1)
        self.assertEqual([run.text for run in runs], ["投标人", "", "名单"])
        self.assertTrue(runs[0].font.bold)
        self.assertTrue(runs[2].font.underline)

    def test_export_word_tables_writes_each_table_to_one_sheet(self):
        with TemporaryDirectory() as tmp_dir:
            source_path = Path(tmp_dir) / "tables.docx"
            output_path = Path(tmp_dir) / "tables.xlsx"

            doc = Document()
            table1 = doc.add_table(rows=1, cols=2)
            table1.cell(0, 0).text = "项目"
            table1.cell(0, 1).text = "预算"
            doc.add_paragraph("between")
            table2 = doc.add_table(rows=1, cols=1)
            table2.cell(0, 0).text = "第二张表"
            doc.save(source_path)

            count = export_word_tables_to_excel(str(source_path), str(output_path))
            workbook = load_workbook(output_path)
            try:
                sheet_names = workbook.sheetnames
                first_value = workbook["表格1"]["A1"].value
                second_value = workbook["表格2"]["A1"].value
            finally:
                workbook.close()

        self.assertEqual(count, 2)
        self.assertEqual(sheet_names, ["表格1", "表格2"])
        self.assertEqual(first_value, "项目")
        self.assertEqual(second_value, "第二张表")

    def test_scan_word_tables_reports_section_context_preview_and_hint(self):
        with TemporaryDirectory() as tmp_dir:
            source_path = Path(tmp_dir) / "tender.docx"

            doc = Document()
            heading = doc.add_paragraph("第五章 评审办法")
            heading.style = "Heading 1"
            doc.add_paragraph("三、商务技术评审")
            doc.add_paragraph("下表为评分标准。")
            table = doc.add_table(rows=2, cols=3)
            table.cell(0, 0).text = "评审因素"
            table.cell(0, 1).text = "分值"
            table.cell(0, 2).text = "评分标准"
            table.cell(1, 0).text = "技术方案"
            table.cell(1, 1).text = "30分"
            table.cell(1, 2).text = "按方案完整性评分"
            doc.save(source_path)

            items = scan_word_tables(str(source_path))

        self.assertEqual(len(items), 1)
        self.assertEqual(items[0].table_index, 1)
        self.assertIn("第五章 评审办法", items[0].section)
        self.assertIn("商务技术评审", items[0].context)
        self.assertIn("评审因素", items[0].preview)
        self.assertIn("评分/评审表", items[0].hint)
        self.assertEqual(items[0].row_count, 2)
        self.assertEqual(items[0].column_count, 3)

    def test_export_word_tables_can_export_selected_table_indexes_only(self):
        with TemporaryDirectory() as tmp_dir:
            source_path = Path(tmp_dir) / "tables.docx"
            output_path = Path(tmp_dir) / "selected.xlsx"

            doc = Document()
            first = doc.add_table(rows=1, cols=1)
            first.cell(0, 0).text = "不要导出"
            second = doc.add_table(rows=1, cols=1)
            second.cell(0, 0).text = "需要导出"
            doc.save(source_path)

            count = export_word_tables_to_excel(
                str(source_path),
                str(output_path),
                table_indexes=[2],
            )
            workbook = load_workbook(output_path)
            try:
                sheet_names = workbook.sheetnames
                value = workbook["表格1"]["A1"].value
            finally:
                workbook.close()

        self.assertEqual(count, 1)
        self.assertEqual(sheet_names, ["表格1"])
        self.assertEqual(value, "需要导出")

    def test_export_word_tables_preserves_horizontal_merge(self):
        with TemporaryDirectory() as tmp_dir:
            source_path = Path(tmp_dir) / "horizontal.docx"
            output_path = Path(tmp_dir) / "horizontal.xlsx"

            doc = Document()
            table = doc.add_table(rows=2, cols=3)
            merged = table.cell(0, 0).merge(table.cell(0, 1))
            merged.text = "横向合并"
            table.cell(0, 2).text = "普通"
            table.cell(1, 0).text = "A"
            table.cell(1, 1).text = "B"
            table.cell(1, 2).text = "C"
            doc.save(source_path)

            export_word_tables_to_excel(str(source_path), str(output_path))
            workbook = load_workbook(output_path)
            try:
                worksheet = workbook["表格1"]
                ranges = {str(cell_range) for cell_range in worksheet.merged_cells.ranges}
                value = worksheet["A1"].value
                tail = worksheet["C1"].value
            finally:
                workbook.close()

        self.assertIn("A1:B1", ranges)
        self.assertEqual(value, "横向合并")
        self.assertEqual(tail, "普通")

    def test_export_word_tables_preserves_vertical_merge(self):
        with TemporaryDirectory() as tmp_dir:
            source_path = Path(tmp_dir) / "vertical.docx"
            output_path = Path(tmp_dir) / "vertical.xlsx"

            doc = Document()
            table = doc.add_table(rows=3, cols=2)
            merged = table.cell(0, 0).merge(table.cell(1, 0))
            merged.text = "纵向合并"
            table.cell(0, 1).text = "右上"
            table.cell(1, 1).text = "右下"
            table.cell(2, 0).text = "末行"
            doc.save(source_path)

            export_word_tables_to_excel(str(source_path), str(output_path))
            workbook = load_workbook(output_path)
            try:
                worksheet = workbook["表格1"]
                ranges = {str(cell_range) for cell_range in worksheet.merged_cells.ranges}
                value = worksheet["A1"].value
                after_merge = worksheet["A3"].value
            finally:
                workbook.close()

        self.assertIn("A1:A2", ranges)
        self.assertEqual(value, "纵向合并")
        self.assertEqual(after_merge, "末行")

    def test_export_word_tables_keeps_paragraph_breaks_in_cells(self):
        with TemporaryDirectory() as tmp_dir:
            source_path = Path(tmp_dir) / "paragraphs.docx"
            output_path = Path(tmp_dir) / "paragraphs.xlsx"

            doc = Document()
            table = doc.add_table(rows=1, cols=1)
            cell = table.cell(0, 0)
            cell.text = "第一段"
            cell.add_paragraph("第二段")
            doc.save(source_path)

            export_word_tables_to_excel(str(source_path), str(output_path))
            workbook = load_workbook(output_path)
            try:
                value = workbook["表格1"]["A1"].value
            finally:
                workbook.close()

        self.assertEqual(value, "第一段\n第二段")

    def test_export_word_tables_skips_document_without_tables(self):
        with TemporaryDirectory() as tmp_dir:
            source_path = Path(tmp_dir) / "plain.docx"
            output_path = Path(tmp_dir) / "plain.xlsx"

            doc = Document()
            doc.add_paragraph("没有表格")
            doc.save(source_path)

            count = export_word_tables_to_excel(str(source_path), str(output_path))

        self.assertEqual(count, 0)
        self.assertFalse(output_path.exists())

    def test_batch_export_word_tables_uses_unique_non_overwriting_output_name(self):
        with TemporaryDirectory() as tmp_dir:
            source_path = Path(tmp_dir) / "sample.docx"
            existing_path = Path(tmp_dir) / "sample_表格.xlsx"
            existing_path.touch()

            doc = Document()
            table = doc.add_table(rows=1, cols=1)
            table.cell(0, 0).text = "内容"
            doc.save(source_path)

            results, skipped, error = batch_export_word_tables([str(source_path)], output_dir=tmp_dir)

        self.assertIsNone(error)
        self.assertEqual(skipped, {})
        self.assertIn("sample.docx", results)
        self.assertEqual(Path(results["sample.docx"]["output_path"]).name, "sample_表格_1.xlsx")
        self.assertEqual(results["sample.docx"]["tables"], 1)

    def test_table_export_output_path_defaults_to_source_directory(self):
        output_path = get_table_export_output_path(r"C:\work\sample.docx")

        self.assertEqual(Path(output_path).name, "sample_表格.xlsx")

    # ---------- 单遍替换：不再跨规则连锁 ----------

    def test_replacement_does_not_cascade_across_rules(self):
        from string_replacer import _prepare_rules, _replace_text_with_rules

        replaced, count = _replace_text_with_rules(
            "采购人名单",
            _prepare_rules([("采购人", "建设单位"), ("单位", "公司")]),
        )

        self.assertEqual(replaced, "建设单位名单")
        self.assertEqual(count, 1)

    def test_simultaneous_rules_replace_original_text_only(self):
        from string_replacer import _prepare_rules, _replace_text_with_rules

        # 原文的 A 和 B 各自按规则替换，B 不会被 A 的结果再改写
        replaced, count = _replace_text_with_rules(
            "AB",
            _prepare_rules([("A", "B"), ("B", "C")]),
        )

        self.assertEqual(replaced, "BC")
        self.assertEqual(count, 2)

    def test_docx_replacement_does_not_cascade_across_rules(self):
        with TemporaryDirectory() as tmp_dir:
            source_path = Path(tmp_dir) / "source.docx"
            output_path = Path(tmp_dir) / "output.docx"

            document = Document()
            document.add_paragraph("采购人名单")
            document.save(source_path)

            count = replace_in_docx(
                str(source_path),
                [("采购人", "建设单位"), ("单位", "公司")],
                str(output_path),
            )

            self.assertEqual(count, 1)
            self.assertEqual(Document(output_path).paragraphs[0].text, "建设单位名单")

    def test_xlsx_replacement_does_not_cascade_across_rules(self):
        with TemporaryDirectory() as tmp_dir:
            source_path = Path(tmp_dir) / "cascade.xlsx"
            workbook = Workbook()
            workbook.active["A1"] = "采购人名单"
            workbook.save(source_path)

            results, error = batch_replace(
                [str(source_path)],
                [("采购人", "建设单位"), ("单位", "公司")],
                output_dir=None,
            )
            output_wb = load_workbook(source_path)
            try:
                value = output_wb.active["A1"].value
            finally:
                output_wb.close()

        self.assertIsNone(error)
        self.assertEqual(results, {"cascade.xlsx": 1})
        self.assertEqual(value, "建设单位名单")

    # ---------- 规则表日期单元格 ----------

    def test_load_replacement_rules_formats_date_cells_without_time_suffix(self):
        with TemporaryDirectory() as tmp_dir:
            rules_path = Path(tmp_dir) / "rules.xlsx"
            workbook = Workbook()
            worksheet = workbook.active
            worksheet["A1"] = "开标日期"
            worksheet["B1"] = datetime.date(2026, 7, 20)
            worksheet["A2"] = "评审时间"
            worksheet["B2"] = datetime.datetime(2026, 7, 20, 9, 30)
            workbook.save(rules_path)

            rules = load_replacement_rules(str(rules_path))

        self.assertEqual(rules, [
            ("开标日期", "2026-07-20"),
            ("评审时间", "2026-07-20 09:30:00"),
        ])

    # ---------- Excel 数值单元格 ----------

    def test_batch_replace_replaces_numeric_excel_cells(self):
        with TemporaryDirectory() as tmp_dir:
            source_path = Path(tmp_dir) / "numbers.xlsx"
            workbook = Workbook()
            workbook.active["A1"] = 20260001
            workbook.save(source_path)

            results, error = batch_replace(
                [str(source_path)],
                [("20260001", "XYZ-2026")],
                output_dir=None,
            )
            output_wb = load_workbook(source_path)
            try:
                value = output_wb.active["A1"].value
            finally:
                output_wb.close()

        self.assertIsNone(error)
        self.assertEqual(results, {"numbers.xlsx": 1})
        self.assertEqual(value, "XYZ-2026")

    def test_batch_replace_keeps_formula_cells_untouched(self):
        with TemporaryDirectory() as tmp_dir:
            source_path = Path(tmp_dir) / "formula.xlsx"
            workbook = Workbook()
            workbook.active["A1"] = "=SUM(B1:B2)"
            workbook.save(source_path)

            results, error = batch_replace(
                [str(source_path)],
                [("SUM", "ADD")],
                output_dir=None,
            )
            output_wb = load_workbook(source_path)
            try:
                value = output_wb.active["A1"].value
            finally:
                output_wb.close()

        self.assertIsNone(error)
        self.assertEqual(results, {"formula.xlsx": 0})
        self.assertEqual(value, "=SUM(B1:B2)")

    # ---------- Word 覆盖范围：超链接 / 嵌套表格 / 文本框 / 内容控件 / 首页页眉 / 脚注 ----------

    def test_docx_replaces_text_inside_hyperlink(self):
        with TemporaryDirectory() as tmp_dir:
            source_path = Path(tmp_dir) / "hyperlink.docx"
            output_path = Path(tmp_dir) / "output.docx"

            document = Document()
            paragraph = document.add_paragraph()
            hyperlink = paragraph._p.makeelement(qn("w:hyperlink"), {qn("r:id"): "rId1"})
            run = hyperlink.makeelement(qn("w:r"), {})
            run_text = run.makeelement(qn("w:t"), {})
            run_text.text = "采购人"
            run.append(run_text)
            hyperlink.append(run)
            paragraph._p.append(hyperlink)
            document.save(source_path)

            count = replace_in_docx(str(source_path), [("采购人", "建设单位")], str(output_path))

            output_xml = Document(output_path).paragraphs[0]._p.xml
            self.assertEqual(count, 1)
            self.assertIn("建设单位", output_xml)
            self.assertNotIn("采购人", output_xml)

    def test_docx_replaces_text_in_nested_table(self):
        with TemporaryDirectory() as tmp_dir:
            source_path = Path(tmp_dir) / "nested.docx"
            output_path = Path(tmp_dir) / "output.docx"

            document = Document()
            outer = document.add_table(rows=1, cols=1)
            inner = outer.cell(0, 0).add_table(rows=1, cols=1)
            inner.cell(0, 0).text = "采购人"
            document.save(source_path)

            count = replace_in_docx(str(source_path), [("采购人", "建设单位")], str(output_path))

            inner_text = (
                Document(output_path).tables[0].cell(0, 0).tables[0].cell(0, 0).text
            )
            self.assertEqual(count, 1)
            self.assertEqual(inner_text, "建设单位")

    def test_docx_replaces_text_in_textbox(self):
        textbox_run = (
            '<w:r xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" '
            'xmlns:v="urn:schemas-microsoft-com:vml">'
            '<w:pict><v:shape style="width:100pt;height:20pt"><v:textbox>'
            '<w:txbxContent><w:p><w:r><w:t>采购人专用章</w:t></w:r></w:p></w:txbxContent>'
            "</v:textbox></v:shape></w:pict></w:r>"
        )
        with TemporaryDirectory() as tmp_dir:
            source_path = Path(tmp_dir) / "textbox.docx"
            output_path = Path(tmp_dir) / "output.docx"

            document = Document()
            document.add_paragraph()._p.append(parse_xml(textbox_run))
            document.save(source_path)

            count = replace_in_docx(str(source_path), [("采购人", "建设单位")], str(output_path))

            output_xml = Document(output_path).element.xml
            self.assertEqual(count, 1)
            self.assertIn("建设单位专用章", output_xml)
            self.assertNotIn("采购人专用章", output_xml)

    def test_docx_replaces_text_inside_content_control(self):
        content_control = (
            '<w:sdt xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
            "<w:sdtContent><w:p><w:r><w:t>采购人在此</w:t></w:r></w:p></w:sdtContent></w:sdt>"
        )
        with TemporaryDirectory() as tmp_dir:
            source_path = Path(tmp_dir) / "sdt.docx"
            output_path = Path(tmp_dir) / "output.docx"

            document = Document()
            document.element.body.insert(0, parse_xml(content_control))
            document.save(source_path)

            count = replace_in_docx(str(source_path), [("采购人", "建设单位")], str(output_path))

            output_xml = Document(output_path).element.xml
            self.assertEqual(count, 1)
            self.assertIn("建设单位在此", output_xml)
            self.assertNotIn("采购人在此", output_xml)

    def test_docx_replaces_first_page_header_text(self):
        with TemporaryDirectory() as tmp_dir:
            source_path = Path(tmp_dir) / "first-page-header.docx"
            output_path = Path(tmp_dir) / "output.docx"

            document = Document()
            section = document.sections[0]
            section.different_first_page_header_footer = True
            section.first_page_header.paragraphs[0].text = "采购人专用页眉"
            document.save(source_path)

            count = replace_in_docx(str(source_path), [("采购人", "建设单位")], str(output_path))

            header_text = (
                Document(output_path).sections[0].first_page_header.paragraphs[0].text
            )
            self.assertEqual(count, 1)
            self.assertEqual(header_text, "建设单位专用页眉")

    def _write_docx_with_footnote(self, path: Path, footnote_text: str) -> None:
        document = Document()
        document.add_paragraph("正文")
        document.save(path)

        with zipfile.ZipFile(path) as archive:
            entries = {name: archive.read(name) for name in archive.namelist()}
        entries["word/footnotes.xml"] = (
            '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
            '<w:footnotes xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
            f'<w:footnote w:id="1"><w:p><w:r><w:t>{footnote_text}</w:t></w:r></w:p></w:footnote>'
            "</w:footnotes>"
        ).encode("utf-8")
        relationships = entries["word/_rels/document.xml.rels"].decode("utf-8")
        entries["word/_rels/document.xml.rels"] = relationships.replace(
            "</Relationships>",
            '<Relationship Id="rIdFnTest" '
            'Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/footnotes" '
            'Target="footnotes.xml"/></Relationships>',
        ).encode("utf-8")
        content_types = entries["[Content_Types].xml"].decode("utf-8")
        entries["[Content_Types].xml"] = content_types.replace(
            "</Types>",
            '<Override PartName="/word/footnotes.xml" '
            'ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.footnotes+xml"/></Types>',
        ).encode("utf-8")
        with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as archive:
            for name, data in entries.items():
                archive.writestr(name, data)

    def test_docx_replaces_footnote_text(self):
        with TemporaryDirectory() as tmp_dir:
            source_path = Path(tmp_dir) / "footnote.docx"
            output_path = Path(tmp_dir) / "output.docx"
            self._write_docx_with_footnote(source_path, "脚注：采购人")

            count = replace_in_docx(str(source_path), [("采购人", "建设单位")], str(output_path))

            with zipfile.ZipFile(output_path) as archive:
                footnotes_xml = archive.read("word/footnotes.xml").decode("utf-8")
            self.assertEqual(count, 1)
            self.assertIn("脚注：建设单位", footnotes_xml)
            Document(output_path)

    # ---------- PPT 备注页 ----------

    def test_pptx_replaces_notes_text(self):
        with TemporaryDirectory() as tmp_dir:
            source_path = Path(tmp_dir) / "notes.pptx"
            output_path = Path(tmp_dir) / "output.pptx"

            presentation = Presentation()
            slide = presentation.slides.add_slide(presentation.slide_layouts[6])
            slide.notes_slide.notes_text_frame.text = "采购人备注"
            presentation.save(source_path)

            count = replace_in_pptx(str(source_path), [("采购人", "建设单位")], str(output_path))

            notes = Presentation(output_path).slides[0].notes_slide.notes_text_frame.text
            self.assertEqual(count, 1)
            self.assertEqual(notes, "建设单位备注")

    # ---------- Word 表格导出：嵌套表格 ----------

    def test_export_word_tables_includes_nested_table_content(self):
        with TemporaryDirectory() as tmp_dir:
            source_path = Path(tmp_dir) / "nested-tables.docx"
            output_path = Path(tmp_dir) / "nested-tables.xlsx"

            document = Document()
            outer = document.add_table(rows=1, cols=1)
            outer.cell(0, 0).text = "外层"
            inner = outer.cell(0, 0).add_table(rows=1, cols=2)
            inner.cell(0, 0).text = "嵌套甲"
            inner.cell(0, 1).text = "嵌套乙"
            document.save(source_path)

            export_word_tables_to_excel(str(source_path), str(output_path))
            workbook = load_workbook(output_path)
            try:
                value = workbook["表格1"]["A1"].value
            finally:
                workbook.close()

        self.assertIn("外层", value)
        self.assertIn("嵌套甲", value)
        self.assertIn("嵌套乙", value)

    # ---------- 招标信息提取准确性 ----------

    def test_extract_keeps_labeled_announcement_date_over_signature_date(self):
        with TemporaryDirectory() as tmp_dir:
            source_path = Path(tmp_dir) / "notice.docx"

            document = Document()
            document.add_paragraph("第一章 投标邀请")
            document.add_paragraph("招标公告日期：2026年7月2日")
            document.add_paragraph("2026年7月1日")
            document.add_paragraph("第二章 投标人须知")
            document.save(source_path)

            info = extract_project_info(str(source_path))

        self.assertEqual(info.get("招标公告日期"), "2026年7月2日")

    def test_extract_skips_junk_tender_item_rows(self):
        with TemporaryDirectory() as tmp_dir:
            source_path = Path(tmp_dir) / "items.docx"

            document = Document()
            table = document.add_table(rows=4, cols=1)
            table.cell(0, 0).text = "标的名称"
            table.cell(1, 0).text = "中控室大屏"
            table.cell(2, 0).text = "备注"
            table.cell(3, 0).text = "合计"
            document.save(source_path)

            info = extract_project_info(str(source_path))

        self.assertEqual(info.get("标的名称"), "中控室大屏")

    def test_extract_drops_reference_only_opening_time_values(self):
        with TemporaryDirectory() as tmp_dir:
            source_path = Path(tmp_dir) / "reference.docx"

            document = Document()
            document.add_paragraph("开标时间：详见第六章")
            document.add_paragraph("开标地点：另行通知")
            document.save(source_path)

            info = extract_project_info(str(source_path))

        self.assertNotIn("开标时间", info)
        self.assertNotIn("开标地点", info)

    def test_extract_keeps_opening_time_with_reference_note_but_real_date(self):
        with TemporaryDirectory() as tmp_dir:
            source_path = Path(tmp_dir) / "with-date.docx"

            document = Document()
            document.add_paragraph("开标时间：2026年7月20日09:30（场地安排详见附件）")
            document.save(source_path)

            info = extract_project_info(str(source_path))

        self.assertIn("2026年7月20日09:30", info.get("开标时间", ""))


if __name__ == "__main__":
    unittest.main()
