import sys
import unittest
from pathlib import Path
from tempfile import TemporaryDirectory
from unittest.mock import patch

from docx import Document
from docx.enum.text import WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches
from lxml import etree
from openpyxl import Workbook
from pptx import Presentation

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from string_replacer import replace_in_docx, replace_in_pptx, replace_in_xlsx


class ReplacementSafetyTests(unittest.TestCase):
    def test_failed_save_preserves_source_and_destination_for_all_new_formats(self):
        with TemporaryDirectory() as directory:
            root = Path(directory)
            doc = Document()
            doc.add_paragraph("OLD")
            workbook = Workbook()
            workbook.active["A1"] = "OLD"
            presentation = Presentation()
            presentation.slides.add_slide(presentation.slide_layouts[1]).shapes.title.text = "OLD"
            cases = (
                (".docx", doc, replace_in_docx, "docx.document.Document.save"),
                (".xlsx", workbook, replace_in_xlsx, "openpyxl.workbook.workbook.Workbook.save"),
                (".xlsm", workbook, replace_in_xlsx, "openpyxl.workbook.workbook.Workbook.save"),
                (".pptx", presentation, replace_in_pptx, "pptx.presentation.Presentation.save"),
            )

            def failed_save(_document, path):
                Path(path).write_bytes(b"partial archive")
                raise OSError("save interrupted")

            for extension, document, replace, save_target in cases:
                for destination in ("source", "existing", "new"):
                    with self.subTest(extension=extension, destination=destination):
                        source = root / ("source" + extension)
                        document.save(source)
                        original = source.read_bytes()
                        output = root / (destination + extension)
                        if destination == "existing":
                            output.write_bytes(b"previous output")
                        previous = output.read_bytes() if output.exists() else None

                        with patch(save_target, failed_save):
                            with self.assertRaisesRegex(OSError, "save interrupted"):
                                replace(str(source), [("OLD", "NEW")], str(output))

                        self.assertEqual(source.read_bytes(), original)
                        if previous is None:
                            self.assertFalse(output.exists())
                        else:
                            self.assertEqual(output.read_bytes(), previous)
                        self.assertEqual(list(root.glob(".replace-simple-*")), [])

    def test_failed_final_replace_preserves_original_and_cleans_temporary_file(self):
        with TemporaryDirectory() as directory:
            source = Path(directory) / "source.docx"
            doc = Document()
            doc.add_paragraph("OLD")
            doc.save(source)
            original = source.read_bytes()
            with patch("os.replace", side_effect=PermissionError("target locked")):
                with self.assertRaisesRegex(PermissionError, "target locked"):
                    replace_in_docx(str(source), [("OLD", "NEW")], str(source))
            self.assertEqual(source.read_bytes(), original)
            self.assertEqual(list(Path(directory).glob(".replace-simple-*")), [])

    def test_docx_preserves_images_and_field_nodes_in_replaced_runs(self):
        with TemporaryDirectory() as directory:
            source = Path(directory) / "source.docx"
            doc = Document()
            paragraph = doc.add_paragraph()
            first = paragraph.add_run("OLD before ")
            first.bold = True
            first.add_picture(str(Path(__file__).resolve().parents[1] / "icon-256.png"), width=Inches(.2))
            first.add_break(WD_BREAK.PAGE)
            first.add_text("OLD after A")
            middle = paragraph.add_run("B")
            middle.add_picture(str(Path(__file__).resolve().parents[1] / "icon-256.png"), width=Inches(.2))
            for tag, value in (("fldChar", "begin"), ("instrText", " PAGE "), ("fldChar", "end")):
                node = OxmlElement("w:" + tag)
                if tag == "instrText":
                    node.text = value
                else:
                    node.set(qn("w:fldCharType"), value)
                middle._r.append(node)
            paragraph.add_run("C tail").italic = True
            doc.save(source)
            preserved = [etree.tostring(node) for node in doc.element.xpath("//w:drawing | //w:fldChar | //w:instrText | //w:br")]

            count = replace_in_docx(str(source), [("OLD", "NEW"), ("ABC", "joined")], str(source))

            result = Document(source)
            self.assertEqual(count, 3)
            self.assertEqual(result.paragraphs[0].text, "NEW before NEW after joined tail")
            self.assertEqual(len(result.inline_shapes), 2)
            self.assertEqual(
                [etree.tostring(node) for node in result.element.xpath("//w:drawing | //w:fldChar | //w:instrText | //w:br")],
                preserved,
            )
            self.assertTrue(result.paragraphs[0].runs[0].bold)
            self.assertTrue(result.paragraphs[0].runs[-1].italic)
            tags = [node.tag for node in result.paragraphs[0].runs[0]._r]
            self.assertLess(tags.index(qn("w:t")), tags.index(qn("w:drawing")))
            self.assertLess(tags.index(qn("w:drawing")), len(tags) - 1)

    def test_docx_replacement_preserves_unmatched_tabs_and_supports_new_whitespace(self):
        with TemporaryDirectory() as directory:
            source = Path(directory) / "source.docx"
            doc = Document()
            run = doc.add_paragraph().add_run("left\tOLD\nOLD right")
            run.add_break(WD_BREAK.PAGE)
            doc.save(source)
            count = replace_in_docx(str(source), [("OLD\nOLD", " new\tline\n ")], str(source))
            result = Document(source)
            self.assertEqual(count, 1)
            self.assertEqual(result.paragraphs[0].text, "left\t new\tline\n  right")
            self.assertEqual(len(result.element.xpath('//w:br[@w:type="page"]')), 1)
            self.assertTrue(result.element.xpath('//w:t[@xml:space="preserve"]'))


if __name__ == "__main__":
    unittest.main()
