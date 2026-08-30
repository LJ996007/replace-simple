import sys
import unittest
from pathlib import Path
from tempfile import TemporaryDirectory

from docx import Document
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls
from openpyxl import load_workbook

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from symbol_clause_extractor import (
    DEFAULT_SECTION_KEYWORDS,
    DEFAULT_SYMBOL_CHARS,
    SYMBOL_CLAUSE_HEADERS,
    SYMBOL_SHEET_TITLE,
    batch_export_symbol_clauses,
    batch_scan_symbol_clause_sections,
    export_symbol_clauses_to_excel,
    extract_symbol_clauses,
    get_symbol_output_path,
    _leading_symbols,
    _split_symbol_clause_blocks,
    normalize_section_keywords,
    normalize_symbol_chars,
    scan_symbol_clause_sections,
    section_matches_keywords,
    strip_clause_symbols,
)
from word_table_exporter import _file_identity


DECIMAL_NUMBERING = (
    '<w:numbering ' + nsdecls("w") + ">"
    '<w:abstractNum w:abstractNumId="100">'
    '<w:lvl w:ilvl="0"><w:start w:val="1"/><w:numFmt w:val="decimal"/>'
    '<w:lvlText w:val="%1、"/></w:lvl>'
    '<w:lvl w:ilvl="1"><w:start w:val="1"/><w:numFmt w:val="chineseCounting"/>'
    '<w:lvlText w:val="%1.%2"/></w:lvl>'
    "</w:abstractNum>"
    '<w:num w:numId="9"><w:abstractNumId w:val="100"/></w:num>'
    "</w:numbering>"
)


def _apply_numbering(paragraph, num_id=9, ilvl=0):
    ppr = paragraph._p.get_or_add_pPr()
    ppr.append(parse_xml(
        f'<w:numPr {nsdecls("w")}><w:ilvl w:val="{ilvl}"/>'
        f'<w:numId w:val="{num_id}"/></w:numPr>'
    ))


def _replace_numbering(document, numbering_xml):
    part = document.part.part_related_by(RT.NUMBERING)
    root = part.element
    for child in list(root):
        root.remove(child)
    for child in parse_xml(numbering_xml):
        root.append(child)


def _make_docx(path, paragraphs=(), table_rows=None, numbered_cell_rows=()):
    """Build a test docx.

    paragraphs: list of (text, num_id_or_None, ilvl) tuples for body paragraphs.
    table_rows: rows of plain cell texts, e.g. [("序号", "名称", "要求"), ...].
    numbered_cell_rows: row indexes whose first-column paragraph uses w:numPr.
    """
    document = Document()
    for text, num_id, ilvl in paragraphs:
        paragraph = document.add_paragraph(text)
        if num_id is not None:
            _apply_numbering(paragraph, num_id=num_id, ilvl=ilvl)

    if table_rows:
        table = document.add_table(rows=len(table_rows), cols=len(table_rows[0]))
        for row_index, row in enumerate(table_rows):
            for column_index, value in enumerate(row):
                cell = table.cell(row_index, column_index)
                cell.paragraphs[0].text = value
                if row_index in numbered_cell_rows and column_index == 0:
                    _apply_numbering(cell.paragraphs[0])

    if paragraphs and any(num_id is not None for _, num_id, _ in paragraphs) or numbered_cell_rows:
        _replace_numbering(document, DECIMAL_NUMBERING)

    document.save(path)
    return path


class LeadingSymbolRuleTests(unittest.TestCase):
    def test_symbol_at_clause_start(self):
        self.assertEqual(_leading_symbols("★分辨率不低于400万像素"), ["★"])
        self.assertEqual(_leading_symbols("  △支持宽动态"), ["△"])

    def test_symbol_after_literal_number(self):
        self.assertEqual(_leading_symbols("3.2、▲支持H.265"), ["▲"])
        self.assertEqual(_leading_symbols("1.★支持夜视"), ["★"])
        self.assertEqual(_leading_symbols("（一）★实质性要求"), ["★"])
        self.assertEqual(_leading_symbols("(3)#重要参数"), ["#"])
        self.assertEqual(_leading_symbols("一、△检测精度"), ["△"])
        self.assertEqual(_leading_symbols("第2条▲必须响应"), ["▲"])
        self.assertEqual(_leading_symbols("１２３★全角序号"), ["★"])

    def test_symbol_after_bracket(self):
        self.assertEqual(_leading_symbols("（★）实质性要求"), ["★"])
        self.assertEqual(_leading_symbols("【▲】关键指标"), ["▲"])

    def test_symbol_only_line_matches(self):
        self.assertEqual(_leading_symbols("★"), ["★"])

    def test_symbol_in_middle_is_ignored(self):
        self.assertEqual(_leading_symbols("该参数为★级，仅供参考"), [])
        self.assertEqual(_leading_symbols("分辨率≥400万像素 ★重要"), [])
        self.assertEqual(_leading_symbols("详见第三章★要求"), [])

    def test_symbol_after_clause_separator_counts(self):
        self.assertEqual(_leading_symbols("★支持A功能；▲支持B功能"), ["★", "▲"])
        self.assertEqual(_leading_symbols("★分辨率≥400万，△支持宽动态"), ["★", "△"])

    def test_multiple_symbols_expand_to_distinct(self):
        self.assertEqual(_leading_symbols("★▲双重要求"), ["★", "▲"])
        self.assertEqual(_leading_symbols("★★强调"), ["★"])

    def test_table_cell_text_splits_at_each_marked_clause(self):
        blocks = _split_symbol_clause_blocks(
            "2.1.1 离子源和进样方式\n"
            "# 2.1.1.1 最大耐受流速：≥2.5 mL/min。\n"
            "（提供彩页或官网证明）\n"
            "# 2.1.1.2 最高加热温度：≥700℃。\n"
            "★2.1.1.3 离子源接口：锥孔结构。\n"
            "2.1.1.4 普通条款不提取\n"
            "★2.1.3 采用180度U型弯曲碰撞室设计。"
        )

        self.assertEqual(blocks, [
            (["#"], "# 2.1.1.1 最大耐受流速：≥2.5 mL/min。\n（提供彩页或官网证明）"),
            (["#"], "# 2.1.1.2 最高加热温度：≥700℃。"),
            (["★"], "★2.1.1.3 离子源接口：锥孔结构。"),
            (["★"], "★2.1.3 采用180度U型弯曲碰撞室设计。"),
        ])


class ExtractBodyParagraphTests(unittest.TestCase):
    def test_body_symbols_extracted_with_full_text(self):
        with TemporaryDirectory() as directory:
            path = _make_docx(Path(directory) / "正文.docx", paragraphs=[
                ("三、采购需求", None, None),
                ("★摄像机分辨率不低于400万像素，具备OSD叠加功能", None, None),
                ("3.2、▲支持H.265编码，码率可调", None, None),
                ("普通条款不提取", None, None),
                ("该参数为★级，仅供参考", None, None),
            ])
            clauses = extract_symbol_clauses(str(path))

        self.assertEqual([clause.symbol for clause in clauses], ["★", "▲"])
        self.assertEqual(clauses[0].text, "★摄像机分辨率不低于400万像素，具备OSD叠加功能")
        self.assertEqual(clauses[1].text, "3.2、▲支持H.265编码，码率可调")
        self.assertIn("采购需求", clauses[0].section)

    def test_auto_numbering_reconstructed(self):
        with TemporaryDirectory() as directory:
            path = _make_docx(Path(directory) / "自动编号.docx", paragraphs=[
                ("采购需求", None, None),
                ("★分辨率不低于400万像素", 9, 0),
                ("△支持OSD叠加", 9, 1),
                ("▲支持宽动态", 9, 1),
                ("★防护等级IP66", 9, 0),
            ])
            clauses = extract_symbol_clauses(str(path))

        self.assertEqual(
            [clause.text for clause in clauses],
            [
                "1、★分辨率不低于400万像素",
                "1.一△支持OSD叠加",
                "1.二▲支持宽动态",
                "2、★防护等级IP66",
            ],
        )

    def test_multiple_symbols_produce_one_row_each(self):
        with TemporaryDirectory() as directory:
            path = _make_docx(Path(directory) / "多符号.docx", paragraphs=[
                ("★支持A功能；▲支持B功能", None, None),
            ])
            clauses = extract_symbol_clauses(str(path))

        self.assertEqual([(c.symbol, c.text) for c in clauses], [
            ("★", "★支持A功能；▲支持B功能"),
            ("▲", "★支持A功能；▲支持B功能"),
        ])


class ExtractTableTests(unittest.TestCase):
    TABLE_ROWS = [
        ("序号", "设备名称", "技术参数要求"),
        ("1", "网络摄像机", "★分辨率不低于400万像素\n▲支持H.265编码"),
        ("2", "交换机", "端口数不少于48个"),
        ("3", "★", "服务器CPU不低于32核"),
    ]

    def test_table_rows_extracted_whole_with_number_column(self):
        with TemporaryDirectory() as directory:
            path = _make_docx(Path(directory) / "表格.docx", table_rows=self.TABLE_ROWS)
            clauses = extract_symbol_clauses(str(path))

        self.assertEqual([clause.symbol for clause in clauses], ["★", "▲", "★"])
        camera_rows = clauses[:2]
        for camera_row in camera_rows:
            self.assertIn("1", camera_row.text.splitlines())
            self.assertIn("网络摄像机", camera_row.text.splitlines())
        self.assertIn("★分辨率不低于400万像素", camera_rows[0].text.splitlines())
        self.assertNotIn("▲支持H.265编码", camera_rows[0].text)
        self.assertIn("▲支持H.265编码", camera_rows[1].text.splitlines())
        self.assertNotIn("★分辨率不低于400万像素", camera_rows[1].text)
        # 无符号行（交换机）不提取
        self.assertNotIn("交换机", "\n".join(clause.text for clause in clauses))
        # 纯符号单元格：整行仍完整提取
        self.assertIn("服务器CPU不低于32核", clauses[2].text)

    def test_multiple_marked_clauses_in_one_cell_export_separately(self):
        with TemporaryDirectory() as directory:
            path = _make_docx(
                Path(directory) / "单元格多条款.docx",
                table_rows=[(
                    "技术要求",
                    "2.1.1 离子源和进样方式\n"
                    "# 2.1.1.1 最大耐受流速：≥2.5 mL/min。\n"
                    "证明材料随附\n"
                    "# 2.1.1.2 最高加热温度：≥700℃。\n"
                    "★2.1.1.3 离子源接口：锥孔结构。\n"
                    "2.1.1.4 普通条款\n"
                    "★2.1.3 采用180度U型弯曲碰撞室设计。",
                )],
            )
            clauses = extract_symbol_clauses(str(path))

        self.assertEqual([clause.symbol for clause in clauses], ["#", "#", "★", "★"])
        self.assertEqual(
            [clause.text for clause in clauses],
            [
                "技术要求\n# 2.1.1.1 最大耐受流速：≥2.5 mL/min。\n证明材料随附",
                "技术要求\n# 2.1.1.2 最高加热温度：≥700℃。",
                "技术要求\n★2.1.1.3 离子源接口：锥孔结构。",
                "技术要求\n★2.1.3 采用180度U型弯曲碰撞室设计。",
            ],
        )

    def test_table_cell_auto_numbering_reconstructed(self):
        with TemporaryDirectory() as directory:
            path = _make_docx(
                Path(directory) / "表格编号.docx",
                table_rows=[
                    ("序号", "设备", "要求"),
                    ("", "摄像机", "支持夜间成像"),
                    ("", "摄像机", "★自动编号条款"),
                ],
                numbered_cell_rows=(1, 2),
            )
            clauses = extract_symbol_clauses(str(path))

        self.assertEqual(len(clauses), 1)
        self.assertEqual(clauses[0].symbol, "★")
        self.assertEqual(clauses[0].text, "2、\n摄像机\n★自动编号条款")


class ExcelExportTests(unittest.TestCase):
    def test_export_writes_three_column_sheet(self):
        with TemporaryDirectory() as directory:
            directory = Path(directory)
            path = _make_docx(directory / "导出.docx", paragraphs=[
                ("★指标一", None, None),
                ("3.2、▲指标二", None, None),
            ])
            output = get_symbol_output_path(str(path), str(directory))
            count = export_symbol_clauses_to_excel(str(path), output)
            self.assertEqual(count, 2)
            self.assertTrue(output.endswith("_指标参数.xlsx"))

            workbook = load_workbook(output)
            worksheet = workbook.active
            self.assertEqual(worksheet.title, SYMBOL_SHEET_TITLE)
            self.assertEqual(
                [worksheet.cell(1, column).value for column in range(1, 4)],
                list(SYMBOL_CLAUSE_HEADERS),
            )
            self.assertEqual([worksheet.cell(2, column).value for column in range(1, 4)],
                             ["1", "★", "★指标一"])
            self.assertEqual([worksheet.cell(3, column).value for column in range(1, 4)],
                             ["2", "▲", "3.2、▲指标二"])

    def test_export_can_strip_symbols_from_content_column(self):
        with TemporaryDirectory() as directory:
            directory = Path(directory)
            path = _make_docx(directory / "去符号导出.docx", paragraphs=[
                ("★指标一", None, None),
                ("3.2、▲指标二", None, None),
            ])
            output = get_symbol_output_path(str(path), str(directory))
            count = export_symbol_clauses_to_excel(
                str(path), output, keep_symbols_in_text=False
            )
            self.assertEqual(count, 2)
            workbook = load_workbook(output)
            worksheet = workbook.active
            self.assertEqual(
                [worksheet.cell(2, column).value for column in range(1, 4)],
                ["1", "★", "指标一"],
            )
            self.assertEqual(
                [worksheet.cell(3, column).value for column in range(1, 4)],
                ["2", "▲", "3.2、指标二"],
            )

    def test_output_path_does_not_overwrite(self):
        with TemporaryDirectory() as directory:
            directory = Path(directory)
            path = _make_docx(directory / "重名.docx", paragraphs=[("★条款", None, None)])
            first = get_symbol_output_path(str(path), str(directory))
            Path(first).write_bytes(b"placeholder")
            second = get_symbol_output_path(str(path), str(directory))
            self.assertNotEqual(first, second)
            self.assertIn("_指标参数_1", second)

    def test_batch_skips_file_without_symbols(self):
        with TemporaryDirectory() as directory:
            directory = Path(directory)
            docx_with = _make_docx(directory / "有符号.docx", paragraphs=[("★条款", None, None)])
            docx_without = _make_docx(directory / "无符号.docx", paragraphs=[("普通条款", None, None)])
            (directory / "说明.txt").write_text("占位", encoding="utf-8")

            progress = []
            results, skipped, error = batch_export_symbol_clauses(
                [str(docx_with), str(docx_without), str(directory / "说明.txt")],
                output_dir=str(directory),
                progress_callback=lambda current, total, name: progress.append(name),
            )

            self.assertIsNone(error)
            self.assertEqual(len(progress), 3)
            self.assertEqual(list(results), ["有符号.docx"])
            self.assertEqual(results["有符号.docx"]["count"], 1)
            self.assertTrue(results["有符号.docx"]["output_path"].endswith("有符号_指标参数.xlsx"))
            self.assertEqual(
                skipped,
                {"无符号.docx": "未找到带符号条款", "说明.txt": "不支持的文件格式"},
            )


class SymbolSectionScanTests(unittest.TestCase):
    def test_scan_groups_clauses_by_exact_section_with_metadata(self):
        with TemporaryDirectory() as directory:
            path = _make_structured_docx(Path(directory) / "分章扫描.docx", [
                ("p", "第五章 采购需求"),
                ("p", "一、摄像设备"),
                ("p", "★分辨率不低于400万像素"),
                ("p", "▲支持H.265编码"),
                ("p", "二、存储设备"),
                ("table", [("1", "存储", "#容量不低于8TB")]),
            ])
            items = scan_symbol_clause_sections(str(path))

        self.assertEqual(len(items), 2)
        self.assertIn("采购需求 / 一、摄像设备", items[0].section)
        self.assertEqual(items[0].symbols, ("★", "▲"))
        self.assertEqual(items[0].clause_count, 2)
        self.assertEqual(items[0].sources, ("正文",))
        self.assertIn("分辨率", items[0].preview)
        self.assertIn("采购需求 / 二、存储设备", items[1].section)
        self.assertEqual(items[1].sources, ("表格1",))

    def test_scan_keeps_unrecognized_section_selectable(self):
        with TemporaryDirectory() as directory:
            path = _make_docx(
                Path(directory) / "无标题.docx",
                paragraphs=[("★未归入标题的条款", None, None)],
            )
            items = scan_symbol_clause_sections(str(path))

        self.assertEqual(len(items), 1)
        self.assertEqual(items[0].section, "未识别章节")

    def test_batch_scan_reports_files_without_symbol_clauses(self):
        with TemporaryDirectory() as directory:
            directory = Path(directory)
            first = _make_docx(directory / "有符号.docx", paragraphs=[("★条款", None, None)])
            second = _make_docx(directory / "无符号.docx", paragraphs=[("普通内容", None, None)])
            items, skipped, error = batch_scan_symbol_clause_sections([str(first), str(second)])

        self.assertIsNone(error)
        self.assertEqual(len(items), 1)
        self.assertEqual(skipped, {"无符号.docx": "未找到带符号条款"})

    def test_batch_scan_reports_corrupt_docx_without_losing_other_results(self):
        with TemporaryDirectory() as directory:
            directory = Path(directory)
            valid = _make_docx(directory / "正常.docx", paragraphs=[("★正常条款", None, None)])
            corrupt = directory / "损坏.docx"
            corrupt.write_bytes(b"not-a-docx")
            items, skipped, error = batch_scan_symbol_clause_sections(
                [str(valid), str(corrupt)]
            )

        self.assertEqual(len(items), 1)
        self.assertEqual(skipped, {})
        self.assertIn("损坏.docx", error)

    def test_batch_export_uses_exact_sections_independently_per_file(self):
        with TemporaryDirectory() as directory:
            directory = Path(directory)
            first = _make_structured_docx(directory / "甲.docx", [
                ("p", "第三章 采购需求"),
                ("p", "★甲需求"),
                ("p", "第四章 评审办法"),
                ("p", "★甲引用"),
            ])
            second = _make_structured_docx(directory / "乙.docx", [
                ("p", "第五章 设备技术规范"),
                ("p", "▲乙规范"),
                ("p", "第六章 响应格式"),
                ("p", "▲乙引用"),
            ])
            selected = {
                _file_identity(str(first)): ["第三章 采购需求"],
                _file_identity(str(second)): ["第五章 设备技术规范"],
            }
            results, skipped, error = batch_export_symbol_clauses(
                [str(first), str(second)],
                output_dir=str(directory),
                section_keywords=["不会命中"],
                selected_sections=selected,
            )

            self.assertIsNone(error)
            self.assertEqual(skipped, {})
            self.assertEqual(set(results), {"甲.docx", "乙.docx"})
            first_book = load_workbook(results["甲.docx"]["output_path"])
            second_book = load_workbook(results["乙.docx"]["output_path"])
            try:
                self.assertEqual(first_book.active["C2"].value, "★甲需求")
                self.assertEqual(second_book.active["C2"].value, "▲乙规范")
                self.assertEqual(first_book.active.max_row, 2)
                self.assertEqual(second_book.active.max_row, 2)
            finally:
                first_book.close()
                second_book.close()

    def test_exact_section_export_handles_duplicate_filenames(self):
        with TemporaryDirectory() as directory:
            directory = Path(directory)
            first_dir = directory / "甲目录"
            second_dir = directory / "乙目录"
            first_dir.mkdir()
            second_dir.mkdir()
            first = _make_structured_docx(first_dir / "同名.docx", [
                ("p", "第三章 采购需求"),
                ("p", "★甲条款"),
            ])
            second = _make_structured_docx(second_dir / "同名.docx", [
                ("p", "第五章 设备技术规范"),
                ("p", "▲乙条款"),
            ])
            selected = {
                _file_identity(str(first)): ["第三章 采购需求"],
                _file_identity(str(second)): ["第五章 设备技术规范"],
            }
            results, skipped, error = batch_export_symbol_clauses(
                [str(first), str(second)],
                output_dir=str(directory),
                selected_sections=selected,
            )

            self.assertIsNone(error)
            self.assertEqual(skipped, {})
            self.assertEqual(len(results), 2)
            output_paths = [info["output_path"] for info in results.values()]
            self.assertEqual(len(set(output_paths)), 2)
            exported_texts = set()
            for output_path in output_paths:
                workbook = load_workbook(output_path)
                try:
                    exported_texts.add(workbook.active["C2"].value)
                finally:
                    workbook.close()
            self.assertEqual(exported_texts, {"★甲条款", "▲乙条款"})


class CustomSymbolSetTests(unittest.TestCase):
    def test_default_symbols_are_the_four_common_markers(self):
        self.assertEqual(DEFAULT_SYMBOL_CHARS, "★#△▲")

    def test_normalize_symbol_chars_filters_and_dedupes(self):
        self.assertEqual(normalize_symbol_chars("★ ▲"), "★▲")
        self.assertEqual(normalize_symbol_chars("★★▲"), "★▲")
        self.assertEqual(normalize_symbol_chars("★1a中 空格▲"), "★▲")
        self.assertEqual(normalize_symbol_chars("◆◇＃＊"), "◆◇＃＊")
        self.assertEqual(normalize_symbol_chars(""), "")
        self.assertEqual(normalize_symbol_chars("12ab一二三"), "")

    def test_default_set_excludes_other_markers(self):
        self.assertEqual(_leading_symbols("☆普通建议条款"), [])
        self.assertEqual(_leading_symbols("◆重要参数"), [])

    def test_extract_with_custom_symbols(self):
        with TemporaryDirectory() as directory:
            path = _make_docx(Path(directory) / "自定义.docx", paragraphs=[
                ("★默认符号条款", None, None),
                ("◆自定义符号条款", None, None),
                ("☆未选符号条款", None, None),
            ])

            only_star = extract_symbol_clauses(str(path), symbols="★")
            self.assertEqual([(c.symbol, c.text) for c in only_star],
                             [("★", "★默认符号条款")])

            diamond_and_star = extract_symbol_clauses(str(path), symbols="★◆☆")
            self.assertEqual([c.symbol for c in diamond_and_star], ["★", "◆", "☆"])

    def test_custom_symbols_work_in_tables_and_numbering(self):
        with TemporaryDirectory() as directory:
            path = _make_docx(
                Path(directory) / "自定义表格.docx",
                paragraphs=[("◆正文条款", 9, 0)],
                table_rows=[("1", "设备", "◆表格条款")],
            )
            clauses = extract_symbol_clauses(str(path), symbols="◆")
            self.assertEqual([c.symbol for c in clauses], ["◆", "◆"])
            self.assertEqual(clauses[0].text, "1、◆正文条款")
            self.assertIn("◆表格条款", clauses[1].text)

    def test_unclean_custom_input_is_sanitized(self):
        with TemporaryDirectory() as directory:
            path = _make_docx(Path(directory) / "清洗.docx", paragraphs=[
                ("★条款一", None, None),
                ("◆条款二", None, None),
            ])
            # 空格、汉字会被清理掉，只剩 ◆
            clauses = extract_symbol_clauses(str(path), symbols=" ◆条 ")
            self.assertEqual([(c.symbol, c.text) for c in clauses],
                             [("◆", "◆条款二")])

    def test_main_default_matches_extractor_default(self):
        from main import DEFAULT_SYMBOL_CHARS as main_default
        self.assertEqual(main_default, DEFAULT_SYMBOL_CHARS)


class StripClauseSymbolTests(unittest.TestCase):
    def test_strips_leading_and_numbered_symbols(self):
        self.assertEqual(strip_clause_symbols("★分辨率不低于400万像素"), "分辨率不低于400万像素")
        self.assertEqual(strip_clause_symbols("3.2、▲支持H.265"), "3.2、支持H.265")
        self.assertEqual(strip_clause_symbols("1.★支持夜视"), "1.支持夜视")
        self.assertEqual(strip_clause_symbols("（★）实质性要求"), "实质性要求")
        self.assertEqual(strip_clause_symbols("【▲】关键指标"), "关键指标")

    def test_strips_multiple_clause_start_symbols_but_keeps_mid_sentence(self):
        self.assertEqual(
            strip_clause_symbols("★支持A功能；▲支持B功能"),
            "支持A功能；支持B功能",
        )
        self.assertEqual(
            strip_clause_symbols("该参数为★级，仅供参考"),
            "该参数为★级，仅供参考",
        )

    def test_extract_can_drop_symbols_from_body_and_table_text(self):
        with TemporaryDirectory() as directory:
            path = _make_docx(
                Path(directory) / "去符号.docx",
                paragraphs=[("3.2、▲支持H.265编码，码率可调", None, None)],
                table_rows=[("1", "网络摄像机", "★分辨率不低于400万像素")],
            )
            clauses = extract_symbol_clauses(str(path), keep_symbols_in_text=False)

        self.assertEqual([c.symbol for c in clauses], ["▲", "★"])
        self.assertEqual(clauses[0].text, "3.2、支持H.265编码，码率可调")
        self.assertNotIn("★", clauses[1].text)
        self.assertIn("分辨率不低于400万像素", clauses[1].text)
        self.assertIn("网络摄像机", clauses[1].text)


class SectionKeywordTests(unittest.TestCase):
    def test_normalize_section_keywords(self):
        self.assertEqual(
            normalize_section_keywords("采购需求，技术要求、 技术规格"),
            ["采购需求", "技术要求", "技术规格"],
        )
        self.assertEqual(normalize_section_keywords([" 采购需求 ", "采购需求", "求"]), ["采购需求"])
        self.assertEqual(normalize_section_keywords(""), [])

    def test_default_keywords_keep_requirement_chapters_and_drop_scoring(self):
        keywords = list(DEFAULT_SECTION_KEYWORDS)
        self.assertTrue(section_matches_keywords("第五章 采购需求", keywords))
        self.assertTrue(section_matches_keywords("第五章 采购需求 / 一、技术规格", keywords))
        self.assertFalse(section_matches_keywords("第六章 评标办法", keywords))
        self.assertFalse(section_matches_keywords("第六章 评标办法 / 技术参数评分", keywords))
        self.assertFalse(section_matches_keywords("第七章 投标文件格式 / 采购需求偏离表", keywords))
        self.assertFalse(section_matches_keywords("未识别章节", keywords))
        self.assertTrue(section_matches_keywords("任何章节", None))

    def test_explicit_scoring_keyword_overrides_exclude(self):
        self.assertTrue(
            section_matches_keywords("第六章 评标办法", ["评标办法"])
        )

    def test_extract_filters_to_requirement_chapters(self):
        with TemporaryDirectory() as directory:
            path = _make_structured_docx(Path(directory) / "分章.docx", [
                ("p", "第五章 采购需求"),
                ("p", "★分辨率不低于400万像素"),
                ("table", [("1", "摄像机", "▲支持H.265")]),
                ("p", "第六章 评标办法"),
                ("p", "★带星号的为实质性要求，见采购需求"),
                ("table", [("★", "评分引用", "见采购需求")]),
            ])
            all_clauses = extract_symbol_clauses(str(path))
            filtered = extract_symbol_clauses(
                str(path),
                section_keywords=list(DEFAULT_SECTION_KEYWORDS),
            )

        self.assertEqual(len(all_clauses), 4)
        self.assertEqual([c.symbol for c in filtered], ["★", "▲"])
        self.assertIn("采购需求", filtered[0].section)
        self.assertNotIn("评标办法", "\n".join(c.section for c in filtered))

    def test_batch_reports_section_miss_separately(self):
        with TemporaryDirectory() as directory:
            directory = Path(directory)
            path = _make_docx(directory / "仅评标.docx", paragraphs=[
                ("第六章 评标办法", None, None),
                ("★实质性要求见采购需求", None, None),
            ])
            results, skipped, error = batch_export_symbol_clauses(
                [str(path)],
                output_dir=str(directory),
                keep_symbols_in_text=False,
                section_keywords=list(DEFAULT_SECTION_KEYWORDS),
            )
            self.assertIsNone(error)
            self.assertEqual(results, {})
            self.assertEqual(skipped, {"仅评标.docx": "未找到指定章节中的带符号条款"})


def _make_structured_docx(path, blocks):
    document = Document()
    for kind, payload in blocks:
        if kind == "p":
            document.add_paragraph(payload)
            continue
        rows = payload
        table = document.add_table(rows=len(rows), cols=len(rows[0]))
        for row_index, row in enumerate(rows):
            for column_index, value in enumerate(row):
                table.cell(row_index, column_index).paragraphs[0].text = value
    document.save(path)
    return path


if __name__ == "__main__":
    unittest.main()
