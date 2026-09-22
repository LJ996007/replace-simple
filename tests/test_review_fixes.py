import os
import time
import unittest
from contextlib import contextmanager
from pathlib import Path
from tempfile import TemporaryDirectory
from types import SimpleNamespace
from unittest.mock import patch

from docx import Document
from openpyxl import Workbook, load_workbook

import app_settings
import legacy_office
import main
import word_scan
from replacement_rules import DELETE_MARKER, rule_conflicts
from string_replacer import (
    batch_replace, load_replacement_rules, replace_in_workbook, save_replacement_rules,
)
from symbol_clause_extractor import SymbolClause, export_symbol_clauses_to_excel
from tender_info_extractor import extract_project_info, extract_project_info_details, extract_project_info_rules
from word_table_exporter import export_word_tables_to_excel


class ReviewFixTests(unittest.TestCase):
    def test_multi_filters_use_original_button_style(self):
        from dataclasses import replace
        from word_table_exporter import TableScanItem
        root = main.tk.Tk()
        root.withdraw()
        try:
            app = main.ReplaceSimpleApp(root, restore_session=False)
            style = main.ttk.Style(root)
            original_styles = {name: (style.configure(name), style.map(name))
                               for name in (".", "TButton", "Accent.TButton", "TCombobox", "Scan.Treeview")}
            app.open_word_table_exporter()
            window = app.table_export_window
            window.file_paths = ["sample.docx"]
            table = TableScanItem("sample.docx", "sample.docx", 1, "第一章", "", "内容", 1, 1, "")
            tables = [replace(table, table_index=i, section=section)
                      for i, section in enumerate(("第一章", "第二章", "第三章"), 1)]
            window._show_scan_result(tables, {}, None, [], {}, None)
            window.select_all_scan_items()
            window.filter_boxes["section"].invoke()
            root.update()
            popup = next(child for child in window.window.winfo_children() if isinstance(child, main.tk.Toplevel))
            self.assertTrue(popup.overrideredirect())

            def descendants(widget):
                for child in widget.winfo_children():
                    yield child
                    yield from descendants(child)

            widgets = list(descendants(popup))
            self.assertFalse(any(isinstance(widget, main.ttk.Button) for widget in widgets))
            chapter = next(
                widget for widget in widgets
                if isinstance(widget, main.tk.Label) and widget.cget("text") == "第三章"
            )
            chapter.event_generate("<Button-1>")
            self.assertEqual(window.scan_filters["section"], {"第一章", "第二章"})
            self.assertEqual(len(window.scan_tree.get_children()), 2)
            self.assertEqual(len(window.selected_scan_keys), 3)
            window._set_multi_filter("type", set())
            self.assertEqual(len(window.scan_tree.get_children()), 0)
            window._set_multi_filter("type", {"表格"})
            self.assertEqual(len(window.scan_tree.get_children()), 2)
            window.reset_scan_filters()
            self.assertEqual(len(window.scan_tree.get_children()), 3)
            self.assertEqual(app.app_bg, "#F0F2F5")
            self.assertEqual(app.surface_bg, "#F7F8FA")
            self.assertEqual(style.lookup("TLabel", "background"), app.surface_bg)
            self.assertEqual(style.lookup("TCheckbutton", "background", ("active",)), app.surface_bg)
            self.assertEqual(style.lookup("TCombobox", "fieldbackground", ("readonly",)), app.surface_bg)
            self.assertEqual(style.lookup("TEntry", "fieldbackground"), app.surface_bg)
            for name, original in original_styles.items():
                self.assertEqual({key: str(value) for key, value in (style.configure(name) or {}).items() if key != "padding"},
                                 {key: str(value) for key, value in (original[0] or {}).items() if key != "padding"})
                self.assertEqual(style.map(name), original[1])
            self.assertEqual(window.filter_boxes["section"].cget("style"), "TButton")
            self.assertEqual(window.export_button.cget("style"), "Accent.TButton")
            more = next(widget for widget in app._busy_widgets if isinstance(widget, main.ttk.Menubutton))
            self.assertEqual(more.cget("text"), "更多 ▼")
            self.assertEqual(more.cget("style"), "TButton")
            menu = more.nametowidget(more.cget("menu"))
            self.assertEqual(menu.entrycget(0, "label"), "将选中规则设为明确删除")
            root.update()
        finally:
            root.destroy()

    def test_indicator_selection_filters_preserve_checks_and_export_only_checked_clauses(self):
        with TemporaryDirectory() as directory:
            root = main.tk.Tk()
            root.withdraw()
            try:
                path = Path(directory) / "selection.docx"
                document = Document()
                document.add_paragraph("★速度要求")
                document.add_paragraph("★精度要求")
                document.add_table(rows=1, cols=1).cell(0, 0).text = "手选内容"
                document.add_table(rows=1, cols=1).cell(0, 0).text = "技术参数"
                document.save(path)
                app = main.ReplaceSimpleApp(root, restore_session=False)
                app._remember_word_files([str(path)])
                app.open_word_table_exporter()
                window = app.table_export_window
                root.update()
                window.window.geometry("1280x900")
                root.update()
                tree_frame = window.scan_tree.master
                detail_frame = window.detail_frame
                self.assertLess(
                    tree_frame.winfo_rootx() + tree_frame.winfo_width(),
                    detail_frame.winfo_rootx() + 4,
                )
                self.assertLess(
                    window.scan_label.winfo_rooty() + window.scan_label.winfo_height(),
                    tree_frame.winfo_rooty(),
                )
                for box in window.filter_boxes.values():
                    self.assertLessEqual(box.winfo_x() + box.winfo_width(), box.master.winfo_width() + 2)
                tables, sections, stamps = word_scan.scan_word_content([str(path)], "★")
                window._show_scan_result(*tables, *sections)
                window.scan_stamps = stamps
                self.assertFalse(window.scan_tree.bind("<Double-1>"))
                window._toggle_scan_iids(["1", "3.1"])
                self.assertEqual(window.scan_tree.item("3", "values")[0], "◩")
                window.scan_tree.selection_set("3.1")
                window._update_scan_detail()
                self.assertIn("速度要求", window.preview_plain_text())
                window.keyword_var.set("精度")
                self.assertEqual(len(window._visible_scan_keys()), 1)
                self.assertIn("另有 2 项已选但被筛掉", window.scan_label.cget("text"))
                window._toggle_scan_iids(["3"])
                self.assertEqual(len(window.selected_scan_keys), 3)
                window.clear_scan_selection()
                self.assertEqual(len(window.selected_scan_keys), 2)
                window.reset_scan_filters()
                window.select_recommended_tables()
                self.assertIn(window._scan_key("table", tables[0][0]), window.selected_scan_keys)
                self.assertEqual(len(window.selected_scan_keys), 3)
                window.only_selected_var.set(True)
                window._apply_scan_filters()
                self.assertEqual(len(window._visible_scan_keys()), 3)
                self.assertEqual(window.export_button.cget("text"), "导出已选 3 项")
                window.output_dir = str(Path(directory) / "out")
                with patch.object(window._task_runner, "submit") as submit:
                    window.start_export()
                result = submit.call_args.args[0](lambda *_: None)
                self.assertIsNone(result[5])
                output = next(iter(result[3].values()))["output_path"]
                workbook = load_workbook(output)
                try:
                    text = " ".join(str(cell.value) for row in workbook.active for cell in row)
                    self.assertIn("速度要求", text)
                    self.assertNotIn("精度要求", text)
                finally:
                    workbook.close()
                window._reset_busy_state()
                window.clear_all_scan_selection()
                self.assertEqual(window.selected_scan_keys, set())
                self.assertEqual(len(window._visible_scan_keys()), 0)
            finally:
                root.update_idletasks()
                root.destroy()

    def test_missing_fields_are_not_deletions_and_conflicts_stop_before_writing(self):
        with TemporaryDirectory() as directory:
            path = Path(directory) / "source.docx"
            document = Document()
            document.add_paragraph("[项目名称] OLD")
            document.save(path)
            rules = extract_project_info_rules(str(path))
            results, error = batch_replace([str(path)], rules)
            self.assertIsNone(error)
            self.assertEqual(sum(results.values()), 0)
            self.assertEqual(Document(path).paragraphs[0].text, "[项目名称] OLD")
            original = path.read_bytes()
            with self.assertRaisesRegex(ValueError, "1、2"):
                batch_replace([str(path)], [("OLD", "one"), ("OLD", "two")])
            self.assertEqual(path.read_bytes(), original)
            batch_replace([str(path)], [("OLD", DELETE_MARKER)])
            self.assertEqual(Document(path).paragraphs[0].text, "[项目名称] ")
            self.assertEqual(rule_conflicts([["", ""], ["A", "one"], ["A", ""], ["A", "two"]]), {"A": [2, 4]})

    def test_rules_round_trip_as_text_and_always_read_first_sheet(self):
        with TemporaryDirectory() as directory:
            path = Path(directory) / "rules.xlsx"
            rows = [("A", "=1+1"), ("B", ""), ("C", DELETE_MARKER)]
            save_replacement_rules(path, rows)
            workbook = load_workbook(path)
            self.assertEqual(workbook.active["B1"].data_type, "s")
            workbook.create_sheet("other").append(["WRONG", "RULE"])
            workbook.active = 1
            workbook.save(path)
            workbook.close()
            self.assertEqual(load_replacement_rules(str(path)), rows)

    def test_excel_text_stays_text_and_original_formulas_are_untouched(self):
        workbook = Workbook()
        sheet = workbook.active
        sheet["A1"] = "OLD"
        sheet["A2"] = "=1+1"
        sheet["A3"] = "=OLD"
        sheet["A3"].data_type = "s"
        count = replace_in_workbook(workbook, [("OLD", "=2+2")])
        self.assertEqual(count, 2)
        self.assertEqual((sheet["A1"].value, sheet["A1"].data_type), ("=2+2", "s"))
        self.assertEqual((sheet["A2"].value, sheet["A2"].data_type), ("=1+1", "f"))
        self.assertEqual((sheet["A3"].value, sheet["A3"].data_type), ("==2+2", "s"))
        cell = SimpleNamespace(Value="OLD", HasFormula=False, NumberFormat="General")
        used = SimpleNamespace(SpecialCells=lambda _kind: SimpleNamespace(Cells=[cell]))
        legacy = SimpleNamespace(Worksheets=[SimpleNamespace(UsedRange=used)])
        self.assertEqual(legacy_office._replace_in_excel_workbook(legacy, [("OLD", "=2+2")]), 1)
        self.assertEqual((cell.Value, cell.NumberFormat), ("=2+2", "@"))

    def test_both_word_exports_preserve_formula_like_text(self):
        with TemporaryDirectory() as directory:
            source = Path(directory) / "source.docx"
            output = Path(directory) / "output.xlsx"
            document = Document()
            document.add_table(rows=1, cols=1).cell(0, 0).text = "=1+1"
            document.save(source)
            export_word_tables_to_excel(str(source), str(output))
            workbook = load_workbook(output)
            self.assertEqual((workbook.active["A1"].value, workbook.active["A1"].data_type), ("=1+1", "s"))
            workbook.close()
            export_symbol_clauses_to_excel(str(source), str(output), [SymbolClause("★", "=1+1")])
            workbook = load_workbook(output)
            self.assertEqual((workbook.active["C2"].value, workbook.active["C2"].data_type), ("=1+1", "s"))
            workbook.close()

    def test_extraction_validates_before_selecting_and_reports_conflicting_sources(self):
        with TemporaryDirectory() as directory:
            path = Path(directory) / "source.docx"
            document = Document()
            for text in ("开标时间：详见第8章", "开标时间：详见第八章", "开标时间：2026年2月31日",
                         "开标时间：2026年9月20日 09:30", "项目名称：甲项目"):
                document.add_paragraph(text)
            table = document.add_table(rows=1, cols=2)
            table.cell(0, 0).text = "项目名称"
            table.cell(0, 1).text = "乙项目"
            document.save(path)
            details = {item["field"]: item for item in extract_project_info_details(str(path))}
            self.assertEqual(details["开标时间"]["value"], "2026年9月20日 09:30")
            self.assertEqual(details["开标时间"]["candidates"][0]["source"], "正文 第4段")
            self.assertEqual(details["开标日期"]["status"], "推断")
            self.assertEqual(details["项目名称"]["status"], "冲突")
            self.assertEqual(details["项目名称"]["value"], "")
            self.assertEqual(len(details["项目名称"]["candidates"]), 2)
            self.assertIn("表格1", details["项目名称"]["candidates"][1]["source"])
            self.assertNotIn("项目名称", extract_project_info(str(path)))
            document.add_paragraph("开标时间：2026年9月21日 09:30")
            document.save(path)
            details = {item["field"]: item for item in extract_project_info_details(str(path))}
            self.assertEqual(details["开标日期"]["status"], "冲突")
            self.assertEqual(details["开标日期"]["value"], "")
            document = Document()
            for text in ("开标时间：2026年9月20日 09:30", "开标日期：2026年9月20日", "开标日期：2026年9月21日"):
                document.add_paragraph(text)
            document.save(path)
            self.assertEqual(dict(extract_project_info_rules(str(path)))["[开标日期]"], "")

    def test_scan_parses_once_exports_snapshot_and_rejects_changed_sources(self):
        with TemporaryDirectory() as directory:
            source = Path(directory) / "source.docx"
            output = Path(directory) / "out"
            document = Document()
            document.add_heading("技术要求", 1)
            document.add_paragraph("★关键参数")
            document.add_table(rows=1, cols=1).cell(0, 0).text = "表格内容"
            document.save(source)
            with patch("word_scan.Document", wraps=Document) as reader:
                tables, sections, stamps = word_scan.scan_word_content([str(source)], "★")
            self.assertEqual(reader.call_count, 1)
            self.assertEqual(len(tables[0]), 1)
            self.assertEqual(len(sections[0]), 1)
            with patch("word_table_exporter.Document", side_effect=AssertionError("不应重复读取")), \
                 patch("symbol_clause_extractor.Document", side_effect=AssertionError("不应重复读取")):
                result = word_scan.export_scanned_content(tables[0], sections[0], stamps, str(output), "★", False)
            self.assertIsNone(result[2])
            self.assertIsNone(result[5])
            workbook = load_workbook(next(iter(result[3].values()))["output_path"])
            self.assertEqual(workbook.active["C2"].value, "关键参数")
            workbook.close()
            stat = source.stat()
            os.utime(source, ns=(stat.st_atime_ns, stat.st_mtime_ns + 1_000_000_000))
            previous_files = list(output.iterdir())
            with self.assertRaisesRegex(ValueError, "重新扫描"):
                word_scan.export_scanned_content(tables[0], sections[0], stamps, str(output), "★", False)
            self.assertEqual(list(output.iterdir()), previous_files)
            source.unlink()
            with self.assertRaisesRegex(ValueError, "重新扫描"):
                word_scan.validate_scan_files([str(source)], stamps)

    def test_legacy_combined_scan_converts_only_once(self):
        with TemporaryDirectory() as directory:
            source = Path(directory) / "legacy.doc"
            readable = Path(directory) / "converted.docx"
            source.write_bytes(b"legacy")
            document = Document()
            document.add_paragraph("★参数")
            document.add_table(rows=1, cols=1).cell(0, 0).text = "内容"
            document.save(readable)

            @contextmanager
            def converted(_path, _session):
                yield str(readable)

            with patch("word_scan.LegacyOfficeSession"), patch("word_scan.temporary_docx_source", side_effect=converted) as convert:
                tables, sections, stamps = word_scan.scan_word_content([str(source)], "★")
            self.assertEqual(convert.call_count, 1)
            self.assertEqual(tables[0][0].file_path, str(source))
            self.assertEqual(sections[0][0].file_path, str(source))
            self.assertEqual(len(stamps), 1)

    def test_settings_preserve_pending_rules_and_survive_failed_commit(self):
        with TemporaryDirectory() as directory, patch.dict(os.environ, {"APPDATA": directory}):
            app_settings.save_settings({"rules": [["A", ""], ["B", DELETE_MARKER]], "presets": ["A"]})
            original = Path(app_settings.settings_path()).read_bytes()
            with patch("file_io.os.replace", side_effect=PermissionError("locked")):
                with self.assertRaises(PermissionError):
                    app_settings.save_settings({"presets": []})
            self.assertEqual(Path(app_settings.settings_path()).read_bytes(), original)
            app_settings.save_settings({"presets": []})
            self.assertEqual(app_settings.load_settings()["rules"], [["A", ""], ["B", DELETE_MARKER]])
            self.assertEqual(app_settings.load_settings()["presets"], [])

    def test_gui_retry_uses_only_failed_files_and_original_rules(self):
        with TemporaryDirectory() as directory:
            root = main.tk.Tk()
            root.withdraw()
            try:
                app = main.ReplaceSimpleApp(root, restore_session=False)
                good, bad = Path(directory) / "good.docx", Path(directory) / "bad.docx"
                document = Document()
                document.add_paragraph("OLD")
                document.save(good)
                bad.write_bytes(b"broken")
                app.replace_files = [str(good), str(bad)]
                app.output_dir = str(Path(directory) / "out")
                app.rules_sheet.set_sheet_data([["OLD", "NEW"], ["pending", ""]])

                def finish():
                    deadline = time.monotonic() + 5
                    while app._task_runner.active and time.monotonic() < deadline:
                        root.update()
                        time.sleep(.01)
                    self.assertFalse(app._task_runner.active)

                with patch.object(app, "_show_result_window") as report:
                    app.start_replace()
                    finish()
                    self.assertIn("失败", report.call_args.args[0])
                    self.assertIn(str(Path(app.output_dir) / good.name), report.call_args.args[0])
                    retry = report.call_args.kwargs["retry"]
                    document.save(bad)
                    app.rules_sheet.set_sheet_data([["OLD", "WRONG"]])
                    retry()
                    finish()
                self.assertEqual(Document(Path(app.output_dir) / bad.name).paragraphs[0].text, "NEW")
                self.assertEqual(sorted(p.name for p in Path(app.output_dir).iterdir()), ["bad.docx", "good.docx"])
                app.rules_sheet.set_sheet_data([["OLD", "one"], ["OLD", "two"]])
                with patch("main.messagebox.showwarning") as warning:
                    app.start_replace()
                self.assertIn("1、2", warning.call_args.args[1])
                self.assertFalse(app._task_runner.active)
            finally:
                root.destroy()

    def test_extraction_window_exports_selection_and_blocks_changed_source(self):
        with TemporaryDirectory() as directory:
            root = main.tk.Tk()
            root.withdraw()
            try:
                source = Path(directory) / "tender.docx"
                document = Document()
                document.add_paragraph("★参数")
                document.add_table(rows=1, cols=1).cell(0, 0).text = "表格内容"
                document.save(source)
                app = main.ReplaceSimpleApp(root, restore_session=False)
                app._remember_word_files([str(source)])
                app.open_word_table_exporter()
                exporter = app.table_export_window
                exporter.output_dir = str(Path(directory) / "out")

                def finish():
                    deadline = time.monotonic() + 5
                    while exporter._task_runner.active and time.monotonic() < deadline:
                        root.update()
                        time.sleep(.01)
                    self.assertFalse(exporter._task_runner.active)

                with patch.object(exporter, "_show_result_window"):
                    exporter.scan_content()
                    finish()
                    self.assertEqual(len(exporter.scan_stamps), 1)
                    exporter.select_all_scan_items()
                    exporter.start_export()
                    finish()
                    self.assertEqual(len(list(Path(exporter.output_dir).iterdir())), 2)
                    stat = source.stat()
                    os.utime(source, ns=(stat.st_atime_ns, stat.st_mtime_ns + 1_000_000_000))
                    with patch("word_export_window.messagebox.showerror") as error, \
                         patch("word_export_window.write_error_log", return_value="test.log"):
                        exporter.start_export()
                        finish()
                    self.assertIn("重新扫描", error.call_args.args[1])
                    self.assertEqual(len(list(Path(exporter.output_dir).iterdir())), 2)
            finally:
                root.destroy()


if __name__ == "__main__":
    unittest.main()
