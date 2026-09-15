"""
Extract common project fields from tender .docx files as replacement rules.
"""

import re
from datetime import date
from typing import Dict, Iterable, List, Optional, Sequence, Tuple

from docx import Document


FIELD_DEFINITIONS: Sequence[Tuple[str, str, Sequence[str]]] = (
    ("项目名称", "[项目名称]", ("项目名称", "采购项目名称", "招标项目名称")),
    ("项目编号", "[项目编号]", ("项目编号", "采购项目编号", "招标编号", "采购编号")),
    ("标的名称", "[标的名称]", ("标的名称", "采购标的名称")),
    (
        "采购人名称",
        "[采购人名称]",
        ("采购人名称", "采购人", "采购单位名称", "采购单位", "招标人名称", "招标人", "建设单位名称", "建设单位"),
    ),
    ("采购人联系人", "[采购人联系人]", ("采购人联系人", "采购联系人", "采购人项目联系人")),
    ("采购人电话", "[采购人电话]", ("采购人电话", "采购人联系电话", "采购人联系方式")),
    (
        "开标时间",
        "[开标时间]",
        (
            "提交投标文件截止时间、开标时间",
            "提交投标文件截止时间和开标时间",
            "提交投标文件截止时间及开标时间",
            "投标截止时间、开标时间",
            "投标截止时间和开标时间",
            "提交投标文件截止时间",
            "递交投标文件截止时间",
            "投标截止时间",
            "开标时间",
            "开启时间",
        ),
    ),
    ("开标日期", "[开标日期]", ("开标日期", "开启日期")),
    ("招标公告日期", "[招标公告日期]", ("招标公告日期", "公告日期")),
    ("开标地点", "[开标地点]", ("开标地点", "开启地点")),
    ("报名人数", "[报名人数]", ()),
    ("采购人地址", "[采购人地址]", ("采购人地址", "采购人联系地址", "采购单位地址")),
)

MAX_VALUE_LENGTH = 500

# 标的名称表格列里不应视为标的的占位行
_JUNK_TENDER_ITEM_VALUES = {"备注", "说明", "合计", "小计", "总计"}
_JUNK_TENDER_ITEM_PREFIXES = ("注：", "注:", "备注", "以上")

# 「详见第八章」「另行通知」这类引用性值不是真实开标信息
_REFERENCE_VALUE_PATTERN = re.compile(r"详见|参见|见第|另行(?:通知|公告)")
PROCUREMENT_CONTEXT_FIELDS: Sequence[Tuple[str, Sequence[str]]] = (
    ("采购人名称", ("名称", "名 称", "单位名称", "机构名称")),
    ("采购人联系人", ("联系人", "联 系 人", "项目联系人")),
    ("采购人电话", ("联系方式", "联系电话", "联系号码", "电话", "电 话")),
    ("采购人地址", ("地址", "地 址", "联系地址")),
)

DATE_PATTERN = re.compile(
    r"(?<!\d)(?:\d{4}\s*年\s*\d{1,2}\s*月\s*\d{1,2}\s*日|\d{4}[-/.]\d{1,2}[-/.]\d{1,2}(?!\d))"
)


def _clean_text(value) -> str:
    if value is None:
        return ""
    text = str(value)
    text = text.replace("\u3000", " ")
    text = re.sub(r"[ \t\r\f\v]+", " ", text)
    text = re.sub(r"\n{2,}", "\n", text)
    return text.strip(" \t\r\n:：;；")


def _normalize_key(value: str) -> str:
    text = _clean_text(value)
    text = re.sub(r"\s+", "", text)
    text = text.strip(":：;；")
    return text


def _key_matches_alias(normalized_key: str, alias: str) -> bool:
    normalized_alias = _normalize_key(alias)
    if normalized_key == normalized_alias:
        return True

    suffix = normalized_key.removeprefix(normalized_alias)
    if not suffix or suffix == normalized_key:
        return False

    return bool(re.fullmatch(r"(?:[（(][^）)]{1,12}[）)])+", suffix))


def _trim_value(value: str) -> str:
    text = _clean_text(value)
    text = re.sub(r"^[：:；;\s]+", "", text)
    text = re.sub(r"\s+", " ", text).strip()
    return text[:MAX_VALUE_LENGTH].strip(" \t\r\n:：;；。")


def _looks_like_key(text: str, aliases: Iterable[str]) -> bool:
    normalized = _normalize_key(text)
    return any(_key_matches_alias(normalized, alias) for alias in aliases)


def _looks_like_any_key(text: str) -> bool:
    normalized = _normalize_key(text)
    if not normalized:
        return False
    for _field_name, _placeholder, aliases in FIELD_DEFINITIONS:
        if any(_key_matches_alias(normalized, alias) for alias in aliases):
            return True
    for _field_name, aliases in PROCUREMENT_CONTEXT_FIELDS:
        if any(_key_matches_alias(normalized, alias) for alias in aliases):
            return True
    return False


def _value_from_labeled_text(text: str, aliases: Iterable[str]) -> Optional[str]:
    line = _clean_text(text)
    if not line:
        return None

    prefix = r"^(?:[（(]?[一二三四五六七八九十\d]+[）)、.．]\s*)?"
    for alias in aliases:
        alias_pattern = r"\s*".join(re.escape(character) for character in alias)
        key_pattern = rf"{alias_pattern}(?:[（(][^）)]{{1,12}}[）)])*"
        patterns = (
            rf"{prefix}{key_pattern}\s*[:：]\s*(.+)$",
            rf"{prefix}{key_pattern}\s+(.+)$",
        )
        for pattern in patterns:
            match = re.search(pattern, line)
            if match:
                value = _trim_value(match.group(1))
                if value and not _looks_like_key(value, aliases):
                    return value
    return None


class _FieldValues(dict):
    """提取值及其候选原文位置；对外仍返回普通字典。"""

    def __init__(self):
        super().__init__()
        self.source = ""
        self.candidates = {}


def _set_if_missing(found: _FieldValues, field_name: str, value: Optional[str], source=None) -> None:
    value = _trim_value(value or "")
    if field_name in ("开标时间", "开标日期", "开标地点") and _is_reference_only_value(value):
        return
    if field_name in ("开标时间", "开标日期", "招标公告日期") and DATE_PATTERN.search(re.sub(r"\s+", "", value)) and not _date_from_text(value):
        return
    if value:
        candidate = {"value": value, "source": source or found.source, "kind": "直接识别"}
        candidates = found.candidates.setdefault(field_name, [])
        if candidate not in candidates:
            candidates.append(candidate)
        found.setdefault(field_name, value)


def _field_name_from_context(text: str) -> Optional[str]:
    normalized = _normalize_key(text)
    if not normalized:
        return None

    procurement_contexts = (
        "采购人信息",
        "采购人联系方式",
        "采购单位信息",
        "招标人信息",
    )
    agency_contexts = (
        "采购代理机构信息",
        "代理机构信息",
        "招标代理机构信息",
        "采购代理信息",
    )
    if any(item in normalized for item in agency_contexts):
        return "agency"
    if any(item in normalized for item in procurement_contexts):
        return "procurement"
    return None


def _extract_procurement_context_value(text: str, found: _FieldValues, source=None) -> bool:
    extracted = False
    for field_name, aliases in PROCUREMENT_CONTEXT_FIELDS:
        value = _value_from_labeled_text(text, aliases)
        if value:
            _set_if_missing(found, field_name, value, source)
            extracted = True
    return extracted


def _extract_tender_item_names_from_table(table, found: Dict[str, str]) -> None:
    """Read every non-empty value below a 标的名称 table header."""
    aliases = next(aliases for name, _placeholder, aliases in FIELD_DEFINITIONS if name == "标的名称")
    rows = list(table.rows)
    for header_row_index, row in enumerate(rows):
        cells = [_clean_text(cell.text) for cell in row.cells]
        for column_index, cell_text in enumerate(cells):
            if not _looks_like_key(cell_text, aliases):
                continue

            names: List[str] = []
            normalized_names = set()
            for data_row in rows[header_row_index + 1:]:
                if column_index >= len(data_row.cells):
                    continue
                value = _trim_value(data_row.cells[column_index].text)
                normalized = _normalize_key(value)
                if not normalized or _looks_like_key(value, aliases):
                    continue
                if (
                    normalized in _JUNK_TENDER_ITEM_VALUES
                    or normalized.startswith(_JUNK_TENDER_ITEM_PREFIXES)
                ):
                    continue
                if normalized in normalized_names:
                    continue
                normalized_names.add(normalized)
                names.append(value)

            if names:
                _set_if_missing(found, "标的名称", "；".join(names))
                return


def _extract_from_table(table, found: _FieldValues, source="正文表格") -> None:
    found.source = source
    _extract_tender_item_names_from_table(table, found)
    active_context: Optional[str] = None
    for row_index, row in enumerate(table.rows, 1):
        cells = [_clean_text(cell.text) for cell in row.cells]
        for index, cell_text in enumerate(cells):
            found.source = f"{source} 第{row_index}行第{index + 1}列"
            if not cell_text:
                continue

            context_field = _field_name_from_context(cell_text)
            if context_field:
                active_context = context_field
                if active_context == "procurement":
                    _extract_procurement_context_value(cell_text, found)
                    for field_name, _placeholder, aliases in FIELD_DEFINITIONS:
                        _set_if_missing(found, field_name, _value_from_labeled_text(cell_text, aliases))
                continue

            if active_context == "procurement":
                _extract_procurement_context_value(cell_text, found)
                for field_name, aliases in PROCUREMENT_CONTEXT_FIELDS:
                    if not _looks_like_key(cell_text, aliases):
                        continue
                    for next_column, next_text in enumerate(cells[index + 1:], index + 2):
                        if not next_text:
                            continue
                        if _looks_like_any_key(next_text):
                            break
                        _set_if_missing(found, field_name, next_text, f"{source} 第{row_index}行第{next_column}列")
                        break

            for field_name, _placeholder, aliases in FIELD_DEFINITIONS:
                if field_name == "标的名称":
                    continue
                value = _value_from_labeled_text(cell_text, aliases)
                if value:
                    _set_if_missing(found, field_name, value)
                    continue

                if not _looks_like_key(cell_text, aliases):
                    continue

                for next_column, next_text in enumerate(cells[index + 1:], index + 2):
                    if not next_text:
                        continue
                    if _looks_like_any_key(next_text):
                        break
                    location = f"{source} 第{row_index}行第{next_column}列"
                    if active_context == "procurement" and _extract_procurement_context_value(next_text, found, location):
                        break
                    else:
                        _set_if_missing(found, field_name, next_text, location)
                        break

        for column_index, cell in enumerate(row.cells, 1):
            for nested_index, nested_table in enumerate(cell.tables, 1):
                _extract_from_table(nested_table, found, f"{source} 第{row_index}行第{column_index}列嵌套表格{nested_index}")


def _extract_from_paragraphs(paragraphs, found: _FieldValues, source="正文") -> None:
    active_context: Optional[str] = None
    in_opening_section = False
    for index, paragraph in enumerate(paragraphs, 1):
        found.source = f"{source} 第{index}段"
        text = _clean_text(paragraph.text)
        if not text:
            continue

        normalized = _normalize_key(text)
        is_opening_heading = (
            "开标时间" in normalized
            and "地点" in normalized
            and ("提交投标文件截止时间" in normalized or "投标截止时间" in normalized)
        )
        if is_opening_heading:
            in_opening_section = True
        elif in_opening_section and re.match(r"^[一二三四五六七八九十]+[、.．]", normalized):
            in_opening_section = False

        if in_opening_section:
            _set_if_missing(found, "开标地点", _value_from_labeled_text(text, ("地点",)))

        context_field = _field_name_from_context(text)
        if context_field:
            active_context = context_field
            if active_context == "procurement":
                _extract_procurement_context_value(text, found)
                for field_name, _placeholder, aliases in FIELD_DEFINITIONS:
                    _set_if_missing(found, field_name, _value_from_labeled_text(text, aliases))
            continue

        if active_context == "procurement":
            _extract_procurement_context_value(text, found)

        for field_name, _placeholder, aliases in FIELD_DEFINITIONS:
            _set_if_missing(found, field_name, _value_from_labeled_text(text, aliases))


def _extract_from_document_part(part, found: _FieldValues, source="正文") -> None:
    _extract_from_paragraphs(part.paragraphs, found, source)
    for index, table in enumerate(part.tables, 1):
        _extract_from_table(table, found, f"{source} 表格{index}")


def _date_from_text(text: str) -> Optional[str]:
    compact = (text or "").translate(str.maketrans("０１２３４５６７８９－／．", "0123456789-/."))
    compact = re.sub(r"\s+", "", compact)
    match = DATE_PATTERN.search(compact)
    if not match:
        return None
    value = match.group(0)
    if "年" in value:
        year, month, day = re.fullmatch(r"(\d{4})年(\d{1,2})月(\d{1,2})日", value).groups()
    else:
        year, month, day = re.split(r"[-/.]", value)
    try:
        parsed = date(int(year), int(month), int(day))
    except ValueError:
        return None
    return f"{parsed.year}年{parsed.month}月{parsed.day}日"


def _clean_opening_time(value: str) -> str:
    value = re.sub(
        r"\s*[（(]\s*北\s*京\s*时\s*间\s*[）)]\s*",
        "",
        value or "",
    )
    return _trim_value(value)


def _split_procurement_contact_and_phone(found: Dict[str, str]) -> None:
    """Split values such as “刘老师，010-82314928” into two fields."""
    phone_value = found.get("采购人电话", "")
    phone_parts = re.split(r"[,，]", phone_value, maxsplit=1)
    if len(phone_parts) == 2 and all(part.strip() for part in phone_parts):
        contact, phone = (_trim_value(part) for part in phone_parts)
        if "采购人联系人" not in found:
            found["采购人联系人"] = contact
        found["采购人电话"] = phone
        return

    contact_value = found.get("采购人联系人", "")
    contact_parts = re.split(r"[,，]", contact_value, maxsplit=1)
    if len(contact_parts) == 2 and all(part.strip() for part in contact_parts):
        contact, phone = (_trim_value(part) for part in contact_parts)
        found["采购人联系人"] = contact
        if "采购人电话" not in found:
            found["采购人电话"] = phone


def _extract_announcement_signature_date(document) -> Optional[Tuple[str, str]]:
    """返回第一章最后一个日期及其段落位置，作为待核对的推断。"""
    in_first_chapter = False
    latest_in_chapter = None
    result = None

    for index, paragraph in enumerate(document.paragraphs, 1):
        text = _clean_text(paragraph.text)
        normalized = _normalize_key(text)
        if not normalized:
            continue

        is_first_chapter = bool(re.search(r"第[一1]章", normalized))
        if is_first_chapter:
            in_first_chapter = True
            latest_in_chapter = None
            continue

        if in_first_chapter and re.search(r"第[二2]章", normalized):
            if latest_in_chapter:
                result = latest_in_chapter
            in_first_chapter = False
            continue

        if in_first_chapter:
            date = _date_from_text(text)
            if date:
                latest_in_chapter = (date, f"正文 第{index}段")

    if in_first_chapter and latest_in_chapter:
        result = latest_in_chapter
    return result


def _is_reference_only_value(value: str) -> bool:
    """章节编号不作为真实日期；保留引用说明之前已有的具体值。"""
    if not _REFERENCE_VALUE_PATTERN.search(value or ""):
        return False
    if _date_from_text(value):
        return False
    prefix = _REFERENCE_VALUE_PATTERN.split(value, maxsplit=1)[0].strip(" （(:：")
    return not prefix


def extract_project_info_details(docx_path: str) -> List[dict]:
    """返回字段、候选来源和识别状态，冲突项留空等待核对。"""
    document = Document(docx_path)
    found = _FieldValues()

    _extract_from_document_part(document, found)
    for index, section in enumerate(document.sections, 1):
        for name, part in (("页眉", section.header), ("页脚", section.footer),
                           ("首页页眉", section.first_page_header), ("首页页脚", section.first_page_footer),
                           ("偶数页页眉", section.even_page_header), ("偶数页页脚", section.even_page_footer)):
            if not part.is_linked_to_previous:
                _extract_from_document_part(part, found, f"第{index}节{name}")

    before_split = dict(found)
    _split_procurement_contact_and_phone(found)

    if "开标时间" in found:
        opening_date = _date_from_text(found["开标时间"])
        if opening_date:
            found.setdefault("开标日期", opening_date)
        found["开标时间"] = _clean_opening_time(found["开标时间"])

    # 落款日期只是兜底推断：按「公告日期」标签提取到的值优先
    signature_date = _extract_announcement_signature_date(document)
    if signature_date:
        found.setdefault("招标公告日期", signature_date[0])

    details = []
    for name, placeholder, _aliases in FIELD_DEFINITIONS:
        candidates = list(found.candidates.get(name, []))
        value = found.get(name, "")
        if value and (not candidates or name in ("采购人电话", "采购人联系人") and value != before_split.get(name)):
            if name == "开标日期":
                sources = found.candidates.get("开标时间", [])
                candidates.extend({"value": _date_from_text(c["value"]),
                                   "source": "由开标时间推导：" + c["source"], "kind": "推断"}
                                  for c in sources if _date_from_text(c["value"]))
                source = ""
            elif name == "招标公告日期":
                source = signature_date[1] + "（第一章末尾日期推断，需核对）"
            else:
                sources = found.candidates.get("采购人电话", []) + found.candidates.get("采购人联系人", [])
                source = "由联系人/电话拆分：" + "；".join(c["source"] for c in sources)
            if source:
                candidates.append({"value": value, "source": source, "kind": "推断"})
        direct_values = {c["value"] for c in candidates if c["kind"] == "直接识别"}
        conflict = len(direct_values) > 1 or not direct_values and len({c["value"] for c in candidates}) > 1
        details.append({"field": name, "placeholder": placeholder,
                        "value": "" if conflict else value,
                        "status": "冲突" if conflict else "缺失" if not value else "推断" if any(c["kind"] == "推断" for c in candidates) else "直接识别",
                        "candidates": candidates})
    return details


def extract_project_info(docx_path: str) -> Dict[str, str]:
    return {item["field"]: item["value"] for item in extract_project_info_details(docx_path) if item["value"]}


def extract_project_info_rules(docx_path: str) -> List[Tuple[str, str]]:
    """Return all requested placeholders in stable display order."""
    return [(item["placeholder"], item["value"]) for item in extract_project_info_details(docx_path)]
