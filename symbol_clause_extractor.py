"""
带符号指标参数提取器（招标文件）。

从 .docx 文件中提取带标记符号（★、#、△、▲ 等）的指标条款：符号可能在
条款最开头，也可能紧跟在条款序号之后；条款可能在正文段落里，也可能在
表格里（整行完整提取）。Word 自动编号（w:numPr）会被重建并拼回条款文本
前，保证导出内容“带序号”。结果写入 3 列 xlsx：数量序号 / 符号 /
详细内容（带序号）。直接用 python-docx + openpyxl 读写，无需安装
Word/WPS/Excel。
"""

import contextlib
import os
import re
from dataclasses import dataclass, replace
from typing import Callable, Dict, List, Optional, Tuple

from docx import Document
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.oxml.ns import qn
from docx.text.paragraph import Paragraph

from word_table_exporter import (
    MAX_COLUMN_WIDTH,
    MIN_COLUMN_WIDTH,
    _detect_heading_level,
    _display_width,
    _extract_table,
    _format_export_error,
    _file_identity,
    _iter_body_blocks,
    _result_key,
    _section_text,
    _shorten,
    _thin_border,
)


SYMBOL_OUTPUT_SUFFIX = "_指标参数"
SYMBOL_SHEET_TITLE = "指标参数"
SYMBOL_CLAUSE_HEADERS = ("数量序号", "符号", "详细内容（带序号）")
MAX_SECTION_TEXT_LENGTH = 80
MAX_SECTION_PREVIEW_LENGTH = 180

# 默认只提取这 4 个符号；需要在「自定义符号」里增删时改这里即可。
DEFAULT_SYMBOL_CHARS = "★#△▲"

# 自定义符号对话框里可勾选的常用符号（前 4 个即默认符号，顺序即显示顺序）。
COMMON_SYMBOL_CHOICES = (
    "★", "#", "△", "▲",
    "☆", "▽", "▼", "◆", "◇", "●", "○", "■", "□", "※", "◎", "＊", "＃",
)

# 默认只抽取这些章节标题里的条款（招标文件里 ★ 条款通常集中在采购需求/技术规格）。
DEFAULT_SECTION_KEYWORDS = (
    "采购需求",
    "技术要求",
    "技术规格",
    "技术参数",
    "服务要求",
    "货物需求",
    "参数要求",
)

# 选择章节对话框里可勾选的常用关键词；前若干项即默认关键词。
COMMON_SECTION_CHOICES = DEFAULT_SECTION_KEYWORDS + (
    "采购标的",
    "商务要求",
)

# 「仅提取指定章节」开启时，默认排除这些标题，避免把评标办法里的引用条款一并抽出。
# 用户把某词主动加进章节关键词后，对应排除项不再生效。
DEFAULT_SECTION_EXCLUDE_KEYWORDS = (
    "评标",
    "评审",
    "评分",
    "投标文件格式",
    "响应文件格式",
)

_CN_NUMERAL = "零一二三四五六七八九"
_CN_DIGIT_CHARS = "一二三四五六七八九十百千零〇两"
_NUM_CHAR = f"[0-9０-９{_CN_DIGIT_CHARS}]"

# 条款序号前缀：覆盖 “（一）”“(1)”“1.”“3.2、”“一、”“第1条” 等常见写法，
# 以及全角数字。符号只有出现在条款开头或这个序号前缀之后才算命中。
_CLAUSE_NUMBER_PREFIX = (
    rf"第\s*{_NUM_CHAR}+\s*[章节条款部分项]?\s*"
    rf"|[（(]\s*{_NUM_CHAR}+(?:\s*[.．、]\s*{_NUM_CHAR}+)*\s*[)）]\s*"
    rf"|{_NUM_CHAR}+(?:\s*[.．]\s*{_NUM_CHAR}+)*\s*[、.．:：)）]?\s*"
)
# 符号必须出现在条款起始位置：行首（可带序号前缀/括号），或者紧跟
# 分号、逗号、句号等条款分隔符之后（同一行内常写多个参数条款）。
# 符号集由调用方传入（默认 ★#△▲，可自定义），按符号集缓存编译结果。
_SYMBOL_POSITION_TEMPLATE = (
    r"(?:^|[；;，,。])\s*(?:" + _CLAUSE_NUMBER_PREFIX + r")?"
    r"[\s\(（\[【{〈「『]?\s*([%s]+)"
)
_symbol_pattern_cache: Dict[str, "re.Pattern"] = {}

# 表格单元格里常把多条参数连续写在一起。遇到不带符号的新编号条款时，
# 它应结束前一条带符号条款，但本身不应被误并入导出内容。这里的规则比
# _CLAUSE_NUMBER_PREFIX 更严格：普通行首数字（如“700℃”）不算新条款。
_NUMBERED_CLAUSE_START = re.compile(
    r"^\s*(?:"
    rf"第\s*{_NUM_CHAR}+\s*[章节条款部分项]?"
    rf"|[（(]\s*{_NUM_CHAR}+(?:\s*[.．、]\s*{_NUM_CHAR}+)*\s*[)）]"
    rf"|{_NUM_CHAR}+(?:\s*[.．]\s*{_NUM_CHAR}+)+"
    rf"|{_NUM_CHAR}+\s*[、.．:：)）]"
    r")"
)


def normalize_symbol_chars(text: str) -> str:
    """清理自定义符号输入：去掉空白、数字、字母和汉字，按顺序去重。

    数字、字母、汉字会和条款序号识别冲突（如自定义“1”会把所有带编号的
    条款都算命中），所以一律不允许。
    """
    result: List[str] = []
    for char in text or "":
        if char.isspace() or char.isalnum():
            continue
        if char not in result:
            result.append(char)
    return "".join(result)


def normalize_section_keywords(values) -> List[str]:
    """清理章节关键词：去空白、按逗号/顿号拆分、去掉过短项、按顺序去重。"""
    if values is None:
        parts: List[str] = []
    elif isinstance(values, str):
        parts = re.split(r"[,，、;；\n/／|]+", values)
    elif isinstance(values, (list, tuple)):
        parts = []
        for item in values:
            if isinstance(item, str):
                parts.extend(re.split(r"[,，、;；\n/／|]+", item))
            elif item is not None:
                parts.append(str(item))
    else:
        parts = [str(values)]

    result: List[str] = []
    seen = set()
    for part in parts:
        text = re.sub(r"\s+", "", part or "")
        if len(text) < 2 or text in seen:
            continue
        result.append(text)
        seen.add(text)
    return result


def section_matches_keywords(
    section: str,
    keywords: Optional[List[str]],
    exclude_keywords: Optional[List[str]] = None,
) -> bool:
    """判断条款所在章节是否应提取。

    keywords 为 None 表示不过滤（全文提取）。
    否则章节标题须包含至少一个关键词；默认再排除评标/评分等章节，
    除非用户把排除词自己写进了关键词。
    """
    if keywords is None:
        return True
    normalized = normalize_section_keywords(keywords)
    if not normalized:
        return False
    text = section or ""
    if not any(keyword in text for keyword in normalized):
        return False
    if exclude_keywords is None:
        excludes = list(DEFAULT_SECTION_EXCLUDE_KEYWORDS)
    else:
        excludes = normalize_section_keywords(exclude_keywords)
    active_excludes = [
        item
        for item in excludes
        if not any(item in keyword or keyword in item for keyword in normalized)
    ]
    return not any(item and item in text for item in active_excludes)


def strip_clause_symbols(text: str, symbols: Optional[str] = None) -> str:
    """去掉条款起始位置的标记符号，保留序号和正文。

    只处理与提取规则相同的位置（行首或序号之后），句子中间出现的符号不动。
    """
    if not text:
        return text
    pattern = _symbol_pattern(symbols)
    empty_wrappers = (
        (r"[（(]\s*[)）]", ""),
        (r"[【\[]\s*[】\]]", ""),
        (r"[{〈「『]\s*[}〉」』]", ""),
    )

    def strip_line(line: str) -> str:
        def replacer(match: "re.Match") -> str:
            full = match.group(0)
            symbol_start = match.start(1) - match.start()
            symbol_end = match.end(1) - match.start()
            return full[:symbol_start] + full[symbol_end:]

        stripped = pattern.sub(replacer, line)
        for wrapper_pattern, replacement in empty_wrappers:
            stripped = re.sub(wrapper_pattern, replacement, stripped)
        return stripped

    return "\n".join(strip_line(line) for line in text.splitlines())


def _symbol_pattern(symbols: Optional[str]) -> "re.Pattern":
    chars = normalize_symbol_chars(symbols) if symbols else DEFAULT_SYMBOL_CHARS
    pattern = _symbol_pattern_cache.get(chars)
    if pattern is None:
        pattern = re.compile(_SYMBOL_POSITION_TEMPLATE % re.escape(chars))
        _symbol_pattern_cache[chars] = pattern
    return pattern


def _leading_symbols(line: str, symbols: Optional[str] = None) -> List[str]:
    """Return marker symbols at clause-start positions of a line, in order."""
    active: List[str] = []
    for match in _symbol_pattern(symbols).finditer(line or ""):
        for symbol in match.group(1):
            if symbol not in active:
                active.append(symbol)
    return active


def _split_symbol_clause_blocks(
    text: str,
    symbols: Optional[str] = None,
) -> List[Tuple[List[str], str]]:
    """Split one table cell into independently marked clause blocks.

    A new marker starts a new block. Unmarked continuation lines stay with the
    active block, while an unmarked numbered clause closes it. Text before the
    first marker is cell context rather than a marked clause and is omitted.
    """
    blocks: List[Tuple[List[str], str]] = []
    active_symbols: List[str] = []
    active_lines: List[str] = []

    def flush() -> None:
        nonlocal active_symbols, active_lines
        block_text = _clean_clause_text("\n".join(active_lines))
        if active_symbols and block_text:
            blocks.append((active_symbols, block_text))
        active_symbols = []
        active_lines = []

    for line in _clean_clause_text(text).splitlines():
        line_symbols = _leading_symbols(line, symbols)
        if line_symbols:
            flush()
            active_symbols = line_symbols
            active_lines = [line]
            continue

        if not active_lines:
            continue
        if _NUMBERED_CLAUSE_START.match(line):
            flush()
            continue
        active_lines.append(line)

    flush()
    return blocks


@dataclass
class SymbolClause:
    symbol: str
    text: str
    section: str = ""
    source: str = ""


@dataclass
class SymbolSectionScanItem:
    """One selectable group of symbol clauses in a document section."""

    file_path: str
    filename: str
    section: str
    symbols: Tuple[str, ...]
    clause_count: int
    sources: Tuple[str, ...]
    preview: str
    clauses: Tuple[SymbolClause, ...] = ()


def scan_symbol_clause_sections(
    docx_path: str,
    symbols: Optional[str] = None,
    document=None,
) -> List[SymbolSectionScanItem]:
    """Scan all matching symbol clauses and group them by exact section path."""
    clauses = extract_symbol_clauses(
        docx_path,
        symbols=symbols,
        keep_symbols_in_text=True,
        section_keywords=None,
        document=document,
    )
    grouped: Dict[str, List[SymbolClause]] = {}
    for clause in clauses:
        grouped.setdefault(clause.section or "未识别章节", []).append(clause)

    filename = os.path.basename(docx_path)
    items: List[SymbolSectionScanItem] = []
    for section, section_clauses in grouped.items():
        active_symbols: List[str] = []
        sources: List[str] = []
        previews: List[str] = []
        for clause in section_clauses:
            if clause.symbol not in active_symbols:
                active_symbols.append(clause.symbol)
            if clause.source and clause.source not in sources:
                sources.append(clause.source)
            preview_text = _clean_clause_text(clause.text)
            if preview_text and preview_text not in previews:
                previews.append(preview_text)

        items.append(SymbolSectionScanItem(
            file_path=docx_path,
            filename=filename,
            section=section,
            symbols=tuple(active_symbols),
            clause_count=len(section_clauses),
            sources=tuple(sources),
            preview=_shorten(" / ".join(previews[:3]), MAX_SECTION_PREVIEW_LENGTH),
            clauses=tuple(section_clauses),
        ))
    return items


def batch_scan_symbol_clause_sections(
    file_paths: List[str],
    progress_callback: Optional[Callable[[int, int, str], None]] = None,
    symbols: Optional[str] = None,
) -> Tuple[List[SymbolSectionScanItem], Dict[str, str], Optional[str]]:
    """Scan symbol-bearing sections from multiple Word files."""
    items: List[SymbolSectionScanItem] = []
    skipped: Dict[str, str] = {}
    errors: List[str] = []
    legacy_session = None

    try:
        for index, file_path in enumerate(file_paths):
            filename = os.path.basename(file_path)
            if progress_callback:
                progress_callback(index + 1, len(file_paths), filename)

            extension = os.path.splitext(file_path)[1].lower()
            if extension not in (".doc", ".docx"):
                skipped[_result_key(skipped, filename, file_path)] = "不支持的文件格式"
                continue

            try:
                if extension == ".doc":
                    from legacy_office import LegacyOfficeSession, temporary_docx_source

                    if legacy_session is None:
                        candidate_session = LegacyOfficeSession()
                        candidate_session.__enter__()
                        legacy_session = candidate_session
                    with temporary_docx_source(file_path, legacy_session) as readable_path:
                        file_items = [
                            replace(item, file_path=file_path, filename=filename)
                            for item in scan_symbol_clause_sections(readable_path, symbols=symbols)
                        ]
                else:
                    file_items = scan_symbol_clause_sections(file_path, symbols=symbols)
                if not file_items:
                    skipped[_result_key(skipped, filename, file_path)] = "未找到带符号条款"
                    continue
                items.extend(file_items)
            except Exception as exc:
                errors.append(f"{filename}: {_format_export_error(exc)}")
    finally:
        if legacy_session is not None:
            legacy_session.close()

    return items, skipped, "\n".join(errors) if errors else None


def get_symbol_output_path(docx_path: str, output_dir: Optional[str] = None) -> str:
    """Return a non-overwriting xlsx path for a symbol clause export."""
    directory = output_dir or os.path.dirname(docx_path)
    stem = os.path.splitext(os.path.basename(docx_path))[0]
    base_name = f"{stem}{SYMBOL_OUTPUT_SUFFIX}.xlsx"
    candidate = os.path.join(directory, base_name)
    index = 1

    while os.path.exists(candidate):
        candidate = os.path.join(directory, f"{stem}{SYMBOL_OUTPUT_SUFFIX}_{index}.xlsx")
        index += 1

    return candidate


def extract_symbol_clauses(
    docx_path: str,
    symbols: Optional[str] = None,
    keep_symbols_in_text: bool = True,
    section_keywords: Optional[List[str]] = None,
    document=None,
) -> List[SymbolClause]:
    """Extract symbol-marked clauses from body paragraphs and tables in order.

    symbols 是参与匹配的符号集合字符串；缺省时只匹配默认的 ★#△▲。
    keep_symbols_in_text 为 False 时，条款正文不再保留标记符号，符号只出现在
    导出的「符号」列。section_keywords 为 None 时全文提取；传入列表则只保留
    章节标题命中这些关键词的条款。
    """
    if document is None:
        document = Document(docx_path)
    tracker = _NumberingTracker(document)
    clauses: List[SymbolClause] = []
    headings: Dict[int, str] = {}
    table_index = 0

    for block in _iter_body_blocks(document):
        if isinstance(block, Paragraph):
            # 空段落也要喂给编号跟踪器：Word 中空段落同样占用自动编号。
            text = tracker.numbered_text(block)

            # 章节识别只用段落本身的文字：自动编号重建出的“一、”“1.”前缀
            # 会让普通列表项被误判成章节标题。
            raw_text = re.sub(r"\s+", " ", block.text or "").strip()
            if raw_text:
                heading_level = _detect_heading_level(block, raw_text)
                if heading_level:
                    headings = {
                        level: heading
                        for level, heading in headings.items()
                        if level < heading_level
                    }
                    headings[heading_level] = _shorten(raw_text, MAX_SECTION_TEXT_LENGTH)

            symbols_found: List[str] = []
            for line in text.splitlines():
                for symbol in _leading_symbols(line, symbols):
                    if symbol not in symbols_found:
                        symbols_found.append(symbol)
            if not symbols_found:
                continue

            section = _section_text(headings)
            if not section_matches_keywords(section, section_keywords):
                continue
            clause_text = _finalize_clause_text(text, symbols, keep_symbols_in_text)
            if not clause_text:
                continue
            for symbol in symbols_found:
                clauses.append(SymbolClause(symbol, clause_text, section, "正文"))
            continue

        table_index += 1
        clauses.extend(
            _extract_table_clauses(
                block,
                tracker,
                headings,
                table_index,
                symbols,
                keep_symbols_in_text=keep_symbols_in_text,
                section_keywords=section_keywords,
            )
        )

    return clauses


def export_symbol_clauses_to_excel(
    docx_path: str,
    output_path: str,
    clauses: Optional[List[SymbolClause]] = None,
    symbols: Optional[str] = None,
    keep_symbols_in_text: bool = True,
    section_keywords: Optional[List[str]] = None,
) -> int:
    """Export symbol clauses of one .docx file to one 3-column .xlsx workbook.

    Returns the number of exported clauses. If no clause is found, no workbook
    is written and 0 is returned.
    """
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font
    from openpyxl.utils import get_column_letter

    if clauses is None:
        clauses = extract_symbol_clauses(
            docx_path,
            symbols=symbols,
            keep_symbols_in_text=keep_symbols_in_text,
            section_keywords=section_keywords,
        )
    if not clauses:
        return 0

    workbook = Workbook()
    try:
        worksheet = workbook.active
        worksheet.title = SYMBOL_SHEET_TITLE
        border = _thin_border()
        wrap_top = Alignment(wrap_text=True, vertical="top")
        center_top = Alignment(wrap_text=True, vertical="top", horizontal="center")
        header_font = Font(bold=True)

        for column, header in enumerate(SYMBOL_CLAUSE_HEADERS, start=1):
            cell = worksheet.cell(row=1, column=column, value=header)
            cell.number_format = "@"
            cell.alignment = center_top
            cell.border = border
            cell.font = header_font

        for row_index, clause in enumerate(clauses, start=2):
            values = (str(row_index - 1), clause.symbol, clause.text)
            for column, value in enumerate(values, start=1):
                cell = worksheet.cell(row=row_index, column=column, value=value)
                cell.data_type = "s"
                cell.number_format = "@"
                cell.alignment = center_top if column < 3 else wrap_top
                cell.border = border

        widths = {1: 10, 2: 8}
        widest = MIN_COLUMN_WIDTH
        for clause in clauses:
            for line in clause.text.splitlines() or [""]:
                widest = max(widest, _display_width(line))
        widths[3] = min(MAX_COLUMN_WIDTH, max(MIN_COLUMN_WIDTH, widest + 2))

        for column, width in widths.items():
            worksheet.column_dimensions[get_column_letter(column)].width = width

        from file_io import atomic_output_path
        with atomic_output_path(output_path) as temporary:
            workbook.save(temporary)
        return len(clauses)
    finally:
        workbook.close()


def batch_export_symbol_clauses(
    file_paths: List[str],
    output_dir: Optional[str] = None,
    progress_callback: Optional[Callable[[int, int, str], None]] = None,
    symbols: Optional[str] = None,
    keep_symbols_in_text: bool = True,
    section_keywords: Optional[List[str]] = None,
    selected_sections: Optional[Dict[str, List[str]]] = None,
) -> Tuple[Dict[str, Dict[str, object]], Dict[str, str], Optional[str]]:
    """Batch export symbol clauses.

    Returns (results, skipped, error_text). results maps original filenames to
    {"count": int, "output_path": str}; skipped maps filenames to a reason.
    selected_sections maps normalized file identities to exact section paths;
    when provided, only those per-file sections are exported and the legacy
    section_keywords filter is ignored.
    """
    results: Dict[str, Dict[str, object]] = {}
    skipped: Dict[str, str] = {}
    errors: List[str] = []
    legacy_session = None

    try:
        for index, file_path in enumerate(file_paths):
            filename = os.path.basename(file_path)
            if progress_callback:
                progress_callback(index + 1, len(file_paths), filename)

            extension = os.path.splitext(file_path)[1].lower()
            if extension not in (".doc", ".docx"):
                skipped[_result_key(skipped, filename, file_path)] = "不支持的文件格式"
                continue

            try:
                exact_sections = None
                if selected_sections is not None:
                    exact_sections = set(selected_sections.get(_file_identity(file_path), []))
                    if not exact_sections:
                        skipped[_result_key(skipped, filename, file_path)] = "未选择符号章节"
                        continue

                readable_path = file_path
                readable_context = contextlib.nullcontext(file_path)
                if extension == ".doc":
                    from legacy_office import LegacyOfficeSession, temporary_docx_source

                    if legacy_session is None:
                        candidate_session = LegacyOfficeSession()
                        candidate_session.__enter__()
                        legacy_session = candidate_session
                    readable_context = temporary_docx_source(file_path, legacy_session)

                with readable_context as readable_path:
                    clauses = extract_symbol_clauses(
                        readable_path,
                        symbols=symbols,
                        keep_symbols_in_text=keep_symbols_in_text,
                        section_keywords=None if exact_sections is not None else section_keywords,
                    )
                if exact_sections is not None:
                    clauses = [clause for clause in clauses if clause.section in exact_sections]
                if not clauses:
                    skipped[_result_key(skipped, filename, file_path)] = (
                        "未找到所选章节中的带符号条款"
                        if exact_sections is not None
                        else "未找到指定章节中的带符号条款"
                        if section_keywords is not None
                        else "未找到带符号条款"
                    )
                    continue

                output_path = get_symbol_output_path(file_path, output_dir)
                count = export_symbol_clauses_to_excel(file_path, output_path, clauses)
                if count == 0:
                    skipped[_result_key(skipped, filename, file_path)] = "未找到带符号条款"
                    continue

                results[_result_key(results, filename, file_path)] = {
                    "count": count,
                    "output_path": output_path,
                }
            except Exception as exc:
                errors.append(f"{filename}: {_format_export_error(exc)}")
    finally:
        if legacy_session is not None:
            legacy_session.close()

    return results, skipped, "\n".join(errors) if errors else None


def _clean_clause_text(text: str) -> str:
    """Normalize whitespace per line while keeping full content and line breaks."""
    text = (text or "").replace("\r\a", "").replace("\a", "").replace("\x07", "")
    lines = [re.sub(r"[ \t　]+", " ", line).strip() for line in text.splitlines()]
    return "\n".join(line for line in lines if line)


def _finalize_clause_text(
    text: str,
    symbols: Optional[str],
    keep_symbols_in_text: bool,
) -> str:
    cleaned = _clean_clause_text(text)
    if keep_symbols_in_text:
        return cleaned
    return _clean_clause_text(strip_clause_symbols(cleaned, symbols))


def _extract_table_clauses(
    table,
    tracker: "_NumberingTracker",
    headings: Dict[int, str],
    table_index: int,
    symbols: Optional[str] = None,
    keep_symbols_in_text: bool = True,
    section_keywords: Optional[List[str]] = None,
) -> List[SymbolClause]:
    """Extract independently marked clause blocks from table rows.

    Plain cells in the same row (such as sequence number and item name) are
    retained as context. When one cell contains multiple marked clauses, each
    clause becomes its own result instead of duplicating the whole cell.
    """
    # 先按文档顺序把表格内（含嵌套表格）所有段落喂给编号跟踪器，
    # 得到“段落 -> 带重建序号文本”的映射，之后提取时只做查找。
    numbered: Dict[object, str] = {}
    for p_el in table._tbl.iter(qn("w:p")):
        numbered[p_el] = tracker.numbered_text(Paragraph(p_el, table))

    def paragraph_text_fn(paragraph: Paragraph) -> str:
        return numbered.get(paragraph._p, paragraph.text)

    exported = _extract_table(table, paragraph_text_fn=paragraph_text_fn)
    section = _section_text(headings)

    rows: Dict[int, List[str]] = {}
    for cell in sorted(exported.cells, key=lambda item: (item.row, item.column)):
        cleaned = _clean_clause_text(cell.text)
        if cleaned:
            rows.setdefault(cell.row, []).append(cleaned)

    clauses: List[SymbolClause] = []
    for row_index in sorted(rows):
        cell_texts = rows[row_index]
        cell_blocks = [
            _split_symbol_clause_blocks(cell_text, symbols)
            for cell_text in cell_texts
        ]
        if not any(cell_blocks):
            continue
        if not section_matches_keywords(section, section_keywords):
            continue

        for marked_cell_index, blocks in enumerate(cell_blocks):
            for block_symbols, block_text in blocks:
                parts: List[str] = []
                for cell_index, cell_text in enumerate(cell_texts):
                    if cell_index == marked_cell_index:
                        parts.append(block_text)
                    elif not cell_blocks[cell_index]:
                        # 序号、名称等不含符号的同排单元格是每条参数的公共上下文。
                        parts.append(cell_text)

                clause_text = _finalize_clause_text(
                    "\n".join(parts), symbols, keep_symbols_in_text
                )
                if not clause_text:
                    continue
                for symbol in block_symbols:
                    clauses.append(SymbolClause(
                        symbol,
                        clause_text,
                        section,
                        f"表格{table_index}",
                    ))

    return clauses


class _NumberingTracker:
    """重建 Word 自动编号（w:numPr），把编号文本拼回段落文本前。

    支持 decimal / 中文编号 / 字母 / 罗马数字 / ①②③ / bullet 等常见
    numFmt，按 (numId, ilvl) 维护计数器，上级层级出现时重置下级计数。
    """

    _MAX_STYLE_CHAIN = 8

    def __init__(self, document):
        self._counters: Dict[Tuple[str, int], int] = {}
        self._nums: Dict[str, Dict[str, object]] = {}
        self._abstracts: Dict[str, Dict[int, object]] = {}
        self._load_numbering(document)

    def numbered_text(self, paragraph: Paragraph) -> str:
        text = paragraph.text or ""
        try:
            numpr = self._effective_numpr(paragraph)
        except Exception:
            numpr = None
        if numpr is None:
            return text

        num_id = self._child_val(numpr, "w:numId")
        if num_id is None or num_id == "0":
            return text
        ilvl = self._child_val(numpr, "w:ilvl") or "0"

        prefix = self._next_number_text(num_id, ilvl)
        if not prefix:
            return text
        return prefix + text

    def _load_numbering(self, document) -> None:
        element = self._numbering_element(document)
        if element is None:
            return

        for abstract in element.findall(qn("w:abstractNum")):
            abstract_id = abstract.get(qn("w:abstractNumId"))
            levels: Dict[int, object] = {}
            for lvl in abstract.findall(qn("w:lvl")):
                try:
                    levels[int(lvl.get(qn("w:ilvl")))] = lvl
                except (TypeError, ValueError):
                    continue
            self._abstracts[abstract_id] = levels

        for num in element.findall(qn("w:num")):
            num_id = num.get(qn("w:numId"))
            abstract_ref = num.find(qn("w:abstractNumId"))
            overrides: Dict[int, int] = {}
            for lvl_override in num.findall(qn("w:lvlOverride")):
                start_override = lvl_override.find(qn("w:startOverride"))
                if start_override is None:
                    continue
                try:
                    overrides[int(lvl_override.get(qn("w:ilvl")))] = int(
                        start_override.get(qn("w:val"))
                    )
                except (TypeError, ValueError):
                    continue
            self._nums[num_id] = {
                "abstract": abstract_ref.get(qn("w:val")) if abstract_ref is not None else None,
                "overrides": overrides,
            }

    @staticmethod
    def _numbering_element(document):
        part = None
        try:
            part = document.part.part_related_by(RT.NUMBERING)
        except Exception:
            part = None
        if part is None:
            try:
                for rel in document.part.rels.values():
                    if rel.reltype == RT.NUMBERING:
                        part = rel.target_part
                        break
            except Exception:
                return None
        if part is None:
            return None
        return getattr(part, "element", None)

    def _effective_numpr(self, paragraph: Paragraph):
        ppr = paragraph._p.find(qn("w:pPr"))
        if ppr is not None:
            numpr = ppr.find(qn("w:numPr"))
            if numpr is not None:
                return numpr

        # 编号也可能挂在段落样式上：沿样式基类链向上找。
        style = None
        try:
            style = paragraph.style
        except Exception:
            style = None
        for _ in range(self._MAX_STYLE_CHAIN):
            if style is None:
                return None
            element = getattr(style, "element", None)
            if element is None:
                return None
            style_ppr = element.find(qn("w:pPr"))
            if style_ppr is not None:
                style_numpr = style_ppr.find(qn("w:numPr"))
                if style_numpr is not None:
                    return style_numpr
            try:
                base = style.base_style
            except Exception:
                return None
            if base is style:
                return None
            style = base
        return None

    def _next_number_text(self, num_id: str, ilvl: str) -> str:
        info = self._nums.get(num_id)
        if not info or info.get("abstract") is None:
            return ""
        levels = self._abstracts.get(info["abstract"]) or {}
        try:
            level_index = int(ilvl)
        except (TypeError, ValueError):
            level_index = 0
        if level_index not in levels and 0 in levels:
            level_index = 0
        lvl = levels.get(level_index)
        if lvl is None:
            return ""

        num_fmt = self._child_val(lvl, "w:numFmt") or "decimal"
        lvl_text = self._child_val(lvl, "w:lvlText")
        if lvl_text is None:
            lvl_text = "%1."

        self._advance_counter(num_id, info, levels, level_index)

        if num_fmt == "none":
            return ""
        if num_fmt == "bullet":
            return re.sub(r"%[1-9]", "", lvl_text)

        def render(level_digit: int) -> str:
            index = level_digit - 1
            value = self._counters.get((num_id, index))
            if value is None:
                value = self._level_start(info, levels, index)
            sub_fmt = self._child_val(levels.get(index), "w:numFmt") or num_fmt
            return _format_number(value, sub_fmt or "decimal")

        return re.sub(r"%([1-9])", lambda m: render(int(m.group(1))), lvl_text)

    def _advance_counter(self, num_id: str, info: Dict[str, object], levels: Dict[int, object], level_index: int) -> None:
        # 上级层级重新出现时，Word 会重置所有更深层级的计数。
        for key in [key for key in self._counters if key[0] == num_id and key[1] > level_index]:
            del self._counters[key]

        key = (num_id, level_index)
        current = self._counters.get(key)
        if current is None:
            self._counters[key] = self._level_start(info, levels, level_index)
        else:
            self._counters[key] = current + 1

    def _level_start(self, info: Dict[str, object], levels: Dict[int, object], level_index: int) -> int:
        override = info.get("overrides", {}).get(level_index)
        if override is not None:
            return override
        start = self._child_val(levels.get(level_index), "w:start")
        try:
            return int(start)
        except (TypeError, ValueError):
            return 1

    @staticmethod
    def _child_val(element, tag: str) -> Optional[str]:
        if element is None:
            return None
        child = element.find(qn(tag))
        if child is None:
            return None
        return child.get(qn("w:val"))


def _format_number(value: int, num_fmt: str) -> str:
    if num_fmt == "zeroDecimal":
        return f"{max(value, 0):02d}"
    if num_fmt == "decimalEnclosedCircle":
        return _circled_number(value)
    if num_fmt in {
        "chineseCounting",
        "chineseCountingThousand",
        "chineseLegalSimplified",
        "ideographDigital",
        "japaneseCounting",
    }:
        return _int_to_chinese(value)
    if num_fmt == "ideographTraditional":
        return _heavenly_stem(value)
    if num_fmt == "upperLetter":
        return _letter_number(value).upper()
    if num_fmt == "lowerLetter":
        return _letter_number(value)
    if num_fmt == "upperRoman":
        return _roman_number(value).upper()
    if num_fmt == "lowerRoman":
        return _roman_number(value)
    return str(value)


def _int_to_chinese(value: int) -> str:
    if value <= 0 or value > 9999:
        return str(value)
    if value < 10:
        return _CN_NUMERAL[value]

    units = ["", "十", "百", "千"]
    digits = [int(digit) for digit in str(value)]
    length = len(digits)
    parts: List[str] = []
    zero_pending = False

    for position, digit in enumerate(digits):
        unit = units[length - 1 - position]
        if digit == 0:
            zero_pending = True
            continue
        if zero_pending and parts:
            parts.append("零")
        zero_pending = False
        if digit == 1 and unit == "十" and not parts:
            parts.append("十")
        else:
            parts.append(_CN_NUMERAL[digit] + unit)

    return "".join(parts) or "零"


def _heavenly_stem(value: int) -> str:
    stems = "甲乙丙丁戊己庚辛壬癸"
    if value <= 0:
        return str(value)
    if value <= 10:
        return stems[value - 1]
    return str(value)


def _letter_number(value: int) -> str:
    if value <= 0:
        return str(value)
    letters: List[str] = []
    while value > 0:
        value, remainder = divmod(value - 1, 26)
        letters.append(chr(ord("a") + remainder))
    return "".join(reversed(letters))


def _roman_number(value: int) -> str:
    if value <= 0 or value >= 4000:
        return str(value)
    table = (
        (1000, "m"), (900, "cm"), (500, "d"), (400, "cd"),
        (100, "c"), (90, "xc"), (50, "l"), (40, "xl"),
        (10, "x"), (9, "ix"), (5, "v"), (4, "iv"), (1, "i"),
    )
    parts: List[str] = []
    for threshold, symbol in table:
        while value >= threshold:
            parts.append(symbol)
            value -= threshold
    return "".join(parts)


def _circled_number(value: int) -> str:
    if 1 <= value <= 20:
        return chr(0x2460 + value - 1)
    if 21 <= value <= 35:
        return chr(0x3251 + value - 21)
    if 36 <= value <= 50:
        return chr(0x32B1 + value - 36)
    return str(value)
