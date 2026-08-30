"""
简化版批量文本替换模块。
支持 Word (.doc/.docx)、Excel (.xls/.xlsx/.xlsm)、PowerPoint (.ppt/.pptx) 文件。
"""

import datetime
import os
import shutil
import warnings
from typing import Callable, Dict, List, Optional, Tuple

warnings.filterwarnings("ignore", message="Data Validation extension is not supported")


_WINDOWS_INVALID_FILENAME_REPLACEMENTS = str.maketrans(
    {
        "<": "＜",
        ">": "＞",
        ":": "：",
        '"': "＂",
        "/": "／",
        "\\": "＼",
        "|": "｜",
        "?": "？",
        "*": "＊",
    }
)
_WINDOWS_RESERVED_FILENAME_STEMS = {
    "CON",
    "PRN",
    "AUX",
    "NUL",
    *(f"COM{index}" for index in range(1, 10)),
    *(f"LPT{index}" for index in range(1, 10)),
}

def _clean_text(value) -> str:
    if value is None:
        return ""
    if isinstance(value, datetime.datetime):
        # Excel 日期单元格：无时间部分时只输出日期，避免 "2026-07-20 00:00:00"
        if (value.hour, value.minute, value.second, value.microsecond) == (0, 0, 0, 0):
            return value.strftime("%Y-%m-%d")
        return value.strftime("%Y-%m-%d %H:%M:%S")
    if isinstance(value, datetime.date):
        return value.strftime("%Y-%m-%d")
    if isinstance(value, datetime.time):
        return value.strftime("%H:%M:%S")
    if isinstance(value, float) and value.is_integer():
        return str(int(value)).strip()
    return str(value).strip()


def load_replacement_rules(excel_path: str) -> List[Tuple[str, str]]:
    """从 Excel 第一张工作表读取两列替换规则。"""
    if os.path.splitext(excel_path)[1].lower() == ".xls":
        return _load_xls_replacement_rules(excel_path)

    from openpyxl import load_workbook

    rules = []
    seen = set()
    wb = load_workbook(excel_path, data_only=False)
    ws = wb.active

    try:
        for old_cell, new_cell in ws.iter_rows(min_row=1, max_col=2):
            old_text = _clean_text(old_cell.value)
            new_text = _clean_text(new_cell.value)
            if not old_text:
                continue

            pair = (old_text, new_text)
            if pair in seen:
                continue

            seen.add(pair)
            rules.append(pair)
    finally:
        wb.close()

    return rules


def _load_xls_replacement_rules(excel_path: str) -> List[Tuple[str, str]]:
    try:
        import xlrd
    except ImportError as exc:
        raise RuntimeError("读取 .xls 规则表需要安装 xlrd：pip install xlrd") from exc

    rules = []
    seen = set()
    wb = xlrd.open_workbook(excel_path)
    ws = wb.sheet_by_index(0)

    def _cell_value(row_idx: int, col_idx: int):
        if col_idx >= ws.ncols:
            return ""
        if ws.cell_type(row_idx, col_idx) == xlrd.XL_CELL_DATE:
            # .xls 日期是序列数浮点，先还原成 datetime 再走 _clean_text
            return xlrd.xldate_as_datetime(ws.cell_value(row_idx, col_idx), wb.datemode)
        return ws.cell_value(row_idx, col_idx)

    for row_idx in range(ws.nrows):
        old_text = _clean_text(_cell_value(row_idx, 0))
        new_text = _clean_text(_cell_value(row_idx, 1))
        if not old_text:
            continue

        pair = (old_text, new_text)
        if pair in seen:
            continue

        seen.add(pair)
        rules.append(pair)

    return rules


def _prepare_rules(rules: List[Tuple[str, str]]) -> List[Dict[str, str]]:
    prepared = []
    for order, (old_text, new_text) in enumerate(rules):
        if not old_text:
            continue
        prepared.append({
            "old": old_text,
            "new": new_text,
            "first_char": old_text[0],
            "order": order,
        })
    return sorted(prepared, key=lambda rule: (-len(rule["old"]), rule["order"]))


def _find_rule_regions(text: str, prepared_rules: List[Dict[str, str]]) -> List[Tuple[int, int, str]]:
    """在原文上单遍扫描，返回 [(start, end, new_text)] 匹配区间。

    同一位置优先命中更长的原文（规则已按长度降序排列）。所有区间都基于
    原文计算，替换结果不会再被任何规则（包括其它规则）二次扫描，
    避免「A→B、B→C」式连锁替换把前一条的结果再改写一遍。
    """
    if not text:
        return []

    present_chars = set(text)
    buckets: Dict[str, List[Dict[str, str]]] = {}
    for rule in prepared_rules:
        if rule["first_char"] in present_chars:
            buckets.setdefault(rule["first_char"], []).append(rule)

    regions: List[Tuple[int, int, str]] = []
    scan = 0
    length = len(text)
    while scan < length:
        bucket = buckets.get(text[scan])
        if bucket is None:
            scan += 1
            continue
        for rule in bucket:
            if not text.startswith(rule["old"], scan):
                continue
            end = scan + len(rule["old"])
            if not _is_occurrence_inside_replacement(text, scan, rule["old"], rule["new"]):
                regions.append((scan, end, rule["new"]))
            scan = end
            break
        else:
            scan += 1
    return regions


def _replace_text_with_rules(text: str, prepared_rules: List[Dict[str, str]]) -> Tuple[str, int]:
    if not text:
        return text, 0

    regions = _find_rule_regions(text, prepared_rules)
    if not regions:
        return text, 0

    parts = []
    cursor = 0
    for start, end, new_text in regions:
        parts.append(text[cursor:start])
        parts.append(new_text)
        cursor = end
    parts.append(text[cursor:])
    return "".join(parts), len(regions)


def _build_run_spans(runs) -> List[Dict[str, object]]:
    spans = []
    position = 0

    for index, run in enumerate(runs):
        text = run.text or ""
        end = position + len(text)
        spans.append({
            "index": index,
            "run": run,
            "start": position,
            "end": end,
        })
        position = end

    return spans


def _find_span_for_position(spans: List[Dict[str, object]], position: int) -> Optional[Dict[str, object]]:
    for span in spans:
        if span["start"] <= position < span["end"]:
            return span
    return None


def _is_occurrence_inside_replacement(text: str, start: int, old_text: str, new_text: str) -> bool:
    """判断当前位置的 old_text 是否已经属于完整 new_text，避免重复追加。

    例如规则 A -> A采购 时，第二次处理 A采购 不应再把前缀 A 替换成 A采购。
    """
    if not old_text or old_text not in new_text:
        return False

    offset = new_text.find(old_text)
    while offset != -1:
        replacement_start = start - offset
        replacement_end = replacement_start + len(new_text)
        if (
            replacement_start >= 0
            and replacement_end <= len(text)
            and text[replacement_start:replacement_end] == new_text
        ):
            return True
        offset = new_text.find(old_text, offset + 1)

    return False


def _replace_range_in_runs(runs, spans: List[Dict[str, object]], start: int, end: int, new_text: str) -> None:
    start_span = _find_span_for_position(spans, start)
    end_span = _find_span_for_position(spans, end - 1)
    if start_span is None or end_span is None:
        return

    start_index = start_span["index"]
    end_index = end_span["index"]
    start_run = start_span["run"]
    end_run = end_span["run"]

    start_text = start_run.text or ""
    prefix = start_text[:start - start_span["start"]]

    if start_index == end_index:
        suffix = start_text[end - start_span["start"]:]
        start_run.text = prefix + new_text + suffix
        return

    end_text = end_run.text or ""
    suffix = end_text[end - end_span["start"]:]
    start_run.text = prefix + new_text

    for index in range(start_index + 1, end_index):
        runs[index].text = ""

    end_run.text = suffix


def _apply_rules_to_runs(runs, prepared_rules: List[Dict[str, str]]) -> int:
    """对一组具有 .text 属性的 run（docx run 元素或 pptx Run 对象）单遍替换。"""
    if not runs:
        return 0

    full_text = "".join(run.text or "" for run in runs)
    regions = _find_rule_regions(full_text, prepared_rules)
    if not regions:
        return 0

    spans = _build_run_spans(runs)
    for start, end, new_text in reversed(regions):
        _replace_range_in_runs(runs, spans, start, end, new_text)
    return len(regions)


def _apply_rules_to_paragraph(paragraph, prepared_rules: List[Dict[str, str]]) -> int:
    return _apply_rules_to_runs(paragraph.runs, prepared_rules)


_MC_NAMESPACE = "http://schemas.openxmlformats.org/markup-compatibility/2006"
_MC_ALTERNATE_CONTENT_TAG = "{%s}AlternateContent" % _MC_NAMESPACE
_MC_FALLBACK_TAG = "{%s}Fallback" % _MC_NAMESPACE


def _has_ancestor_tag(element, tag: str) -> bool:
    parent = element.getparent()
    while parent is not None:
        if parent.tag == tag:
            return True
        parent = parent.getparent()
    return False


def _docx_paragraph_runs(p_element) -> List[object]:
    """按文档顺序收集段落中的 run 元素。

    覆盖超链接、智能标记、修订（w:ins）和行内内容控件里的 run；
    不进入 run 元素内部，文本框等嵌套结构由段落级遍历单独处理。
    """
    from docx.oxml.ns import qn

    run_tag = qn("w:r")
    container_tags = {
        qn("w:hyperlink"),
        qn("w:smartTag"),
        qn("w:ins"),
        qn("w:sdt"),
        qn("w:sdtContent"),
        qn("w:fldSimple"),
    }
    runs: List[object] = []

    def walk(element) -> None:
        for child in element:
            if child.tag == run_tag:
                runs.append(child)
            elif child.tag in container_tags:
                walk(child)

    walk(p_element)
    return runs


def _iter_docx_paragraphs(element):
    """产出 XML 子树内的全部段落，含表格单元格、内容控件和文本框。

    AlternateContent 的 Fallback 分支与 Choice 分支内容相同，跳过
    Fallback，避免同一段文字被处理两次。
    """
    from docx.oxml.ns import qn

    has_alternate_content = element.find(f".//{_MC_ALTERNATE_CONTENT_TAG}") is not None
    for p_element in element.iter(qn("w:p")):
        if has_alternate_content and _has_ancestor_tag(p_element, _MC_FALLBACK_TAG):
            continue
        yield p_element


def _apply_rules_to_docx_element(element, prepared_rules: List[Dict[str, str]]) -> int:
    total_count = 0
    for p_element in _iter_docx_paragraphs(element):
        total_count += _apply_rules_to_runs(_docx_paragraph_runs(p_element), prepared_rules)
    return total_count


def _apply_rules_to_docx_headers_footers(document, prepared_rules: List[Dict[str, str]]) -> int:
    """处理各节的默认、首页和偶数页页眉页脚。

    is_linked_to_previous 为 True 时没有自己的定义（内容继承上一节，
    上一节会单独处理），跳过以免重复替换或凭空创建空页眉。
    """
    total_count = 0
    for section in document.sections:
        for container in (
            section.header,
            section.footer,
            section.first_page_header,
            section.first_page_footer,
            section.even_page_header,
            section.even_page_footer,
        ):
            if container.is_linked_to_previous:
                continue
            total_count += _apply_rules_to_docx_element(container._element, prepared_rules)
    return total_count


_DOCX_AUX_PART_SUFFIXES = ("/footnotes.xml", "/endnotes.xml", "/comments.xml")


def _apply_rules_to_docx_aux_parts(document, prepared_rules: List[Dict[str, str]]) -> int:
    """处理脚注、尾注和批注：python-docx 没有对应 API，直接解析 part 内容。"""
    from lxml import etree

    from docx.oxml import parse_xml

    total_count = 0
    for part in document.part.package.iter_parts():
        if not str(part.partname).endswith(_DOCX_AUX_PART_SUFFIXES):
            continue

        element = getattr(part, "element", None)
        if element is not None:
            total_count += _apply_rules_to_docx_element(element, prepared_rules)
            continue

        blob = getattr(part, "blob", None)
        if not blob:
            continue
        root = parse_xml(blob)
        replaced = _apply_rules_to_docx_element(root, prepared_rules)
        if replaced:
            part._blob = etree.tostring(
                root, xml_declaration=True, encoding="UTF-8", standalone=True
            )
        total_count += replaced
    return total_count


def replace_in_docx(file_path: str, rules: List[Tuple[str, str]], output_path: str) -> int:
    from docx import Document

    doc = Document(file_path)
    prepared_rules = _prepare_rules(rules)

    # 直接遍历 body XML 子树：正文、表格（含嵌套）、内容控件、文本框一次覆盖
    total_count = _apply_rules_to_docx_element(doc.element.body, prepared_rules)
    total_count += _apply_rules_to_docx_headers_footers(doc, prepared_rules)
    total_count += _apply_rules_to_docx_aux_parts(doc, prepared_rules)

    doc.save(output_path)
    return total_count


def _workbook_cell_text(value) -> Optional[str]:
    """把单元格值转成可参与文本替换的字符串；不适合替换的返回 None。

    公式格跳过（改文本会破坏公式）；日期/时间格跳过（显示格式由
    number_format 决定，直接替换底层值会改变显示内容）；数值格按
    Excel 的常规显示转成文本，规则原文与其显示内容一致时也能替换。
    """
    if value is None or isinstance(value, bool):
        return None
    if isinstance(value, str):
        return None if value.startswith("=") else value
    if isinstance(value, int):
        return str(value)
    if isinstance(value, float):
        if value != value or value in (float("inf"), float("-inf")):
            return None
        return str(int(value)) if value.is_integer() else str(value)
    return None


def replace_in_workbook(workbook, rules: List[Tuple[str, str]]) -> int:
    prepared_rules = _prepare_rules(rules)
    total_count = 0

    for worksheet in workbook.worksheets:
        for row in worksheet.iter_rows():
            for cell in row:
                text = _workbook_cell_text(cell.value)
                if not text:
                    continue

                new_value, replaced_count = _replace_text_with_rules(text, prepared_rules)
                if replaced_count:
                    cell.value = new_value
                    total_count += replaced_count

    return total_count


def replace_in_xlsx(file_path: str, rules: List[Tuple[str, str]], output_path: str) -> int:
    from openpyxl import load_workbook

    if os.path.abspath(file_path) != os.path.abspath(output_path):
        shutil.copy2(file_path, output_path)

    keep_vba = os.path.splitext(output_path)[1].lower() == ".xlsm"
    workbook = load_workbook(output_path, keep_vba=keep_vba)
    try:
        total_count = replace_in_workbook(workbook, rules)
        workbook.save(output_path)
    finally:
        workbook.close()

    return total_count


def _is_group_shape(shape) -> bool:
    return hasattr(shape, "shapes") and "GROUP" in str(getattr(shape, "shape_type", ""))


def _replace_in_ppt_text_frame(text_frame, prepared_rules: List[Dict[str, str]]) -> int:
    total_count = 0
    for paragraph in text_frame.paragraphs:
        total_count += _apply_rules_to_paragraph(paragraph, prepared_rules)
    return total_count


def _visit_ppt_shape(shape, text_frame_handler: Callable) -> int:
    total_count = 0

    if getattr(shape, "has_text_frame", False):
        total_count += text_frame_handler(shape.text_frame)

    if getattr(shape, "has_table", False):
        for row in shape.table.rows:
            for cell in row.cells:
                if cell.text_frame:
                    total_count += text_frame_handler(cell.text_frame)

    if _is_group_shape(shape):
        for child_shape in shape.shapes:
            total_count += _visit_ppt_shape(child_shape, text_frame_handler)

    return total_count


def replace_in_pptx(file_path: str, rules: List[Tuple[str, str]], output_path: str) -> int:
    from pptx import Presentation

    presentation = Presentation(file_path)
    prepared_rules = _prepare_rules(rules)
    total_count = 0

    for slide in presentation.slides:
        for shape in slide.shapes:
            total_count += _visit_ppt_shape(
                shape,
                lambda text_frame: _replace_in_ppt_text_frame(text_frame, prepared_rules),
            )
        if slide.has_notes_slide:
            total_count += _replace_in_ppt_text_frame(
                slide.notes_slide.notes_text_frame, prepared_rules
            )

    presentation.save(output_path)
    return total_count


def sanitize_windows_filename_stem(stem: str) -> str:
    """把文件名主体转换为可在 Windows 上安全保存的形式。"""
    sanitized = stem.translate(_WINDOWS_INVALID_FILENAME_REPLACEMENTS)
    sanitized = "".join("＿" if ord(char) < 32 else char for char in sanitized)
    sanitized = sanitized.rstrip(" .")
    if not sanitized:
        return "_"

    device_stem = sanitized.split(".", 1)[0].upper()
    if device_stem in _WINDOWS_RESERVED_FILENAME_STEMS:
        sanitized = f"_{sanitized}"
    return sanitized


def apply_rules_to_filename(filename: str, rules: List[Tuple[str, str]]) -> str:
    """替换文件名主体并清理 Windows 禁用字符，保留原扩展名不变。"""
    stem, ext = os.path.splitext(filename)
    replaced_stem, _count = _replace_text_with_rules(stem, _prepare_rules(rules))
    return sanitize_windows_filename_stem(replaced_stem) + ext


def _with_output_index(path: str, index: int) -> str:
    directory, filename = os.path.split(path)
    stem, ext = os.path.splitext(filename)
    return os.path.join(directory, f"{stem}_{index}{ext}")


def _same_file_path(first: str, second: str) -> bool:
    return os.path.normcase(os.path.abspath(os.path.realpath(first))) == os.path.normcase(
        os.path.abspath(os.path.realpath(second))
    )


def _file_identity(path: str) -> str:
    return os.path.normcase(os.path.abspath(os.path.realpath(os.fspath(path))))


def get_output_path(
    file_path: str,
    rules: Optional[List[Tuple[str, str]]] = None,
    output_dir: Optional[str] = None,
    reserved_output_paths: Optional[set] = None,
) -> str:
    """生成输出路径。

    - 文件名按规则替换后若与源文件相同，直接覆盖源文件（不再加「_已替换」）。
    - 若目标路径已被其他已有文件占用，或本批任务中其他文件已占用，则追加 _1、_2…
    """
    dir_path = output_dir or os.path.dirname(file_path)
    if output_dir:
        os.makedirs(output_dir, exist_ok=True)

    filename = os.path.basename(file_path)
    if rules:
        filename = apply_rules_to_filename(filename, rules)

    output_path = os.path.join(dir_path, filename)
    reserved = reserved_output_paths if reserved_output_paths is not None else set()
    candidate = output_path
    index = 1
    while _file_identity(candidate) in reserved or (
        os.path.exists(candidate) and not _same_file_path(file_path, candidate)
    ):
        candidate = _with_output_index(output_path, index)
        index += 1

    reserved.add(_file_identity(candidate))
    return candidate


def _result_key(results: Dict[str, int], filename: str, file_path: str) -> str:
    if filename not in results:
        return filename

    parent = os.path.dirname(file_path)
    candidate = f"{filename} ({parent})"
    index = 2
    while candidate in results:
        candidate = f"{filename} ({parent}, {index})"
        index += 1
    return candidate


def _format_processing_error(exc: Exception) -> str:
    if isinstance(exc, PermissionError):
        return "文件无法读写，请先关闭正在打开的 Word/Excel/PPT 文件后重试"
    return str(exc)


def batch_replace(
    file_paths: List[str],
    rules: List[Tuple[str, str]],
    output_dir: Optional[str] = None,
    progress_callback: Optional[Callable[[int, int, str], None]] = None,
) -> Tuple[Dict[str, int], Optional[str]]:
    """批量处理 Office 文件，返回替换结果和错误信息。"""
    results = {}
    errors = []
    reserved_output_paths = set()
    legacy_session = None

    try:
        for index, file_path in enumerate(file_paths):
            filename = os.path.basename(file_path)
            if progress_callback:
                progress_callback(index + 1, len(file_paths), filename)

            try:
                output_path = get_output_path(
                    file_path,
                    rules,
                    output_dir,
                    reserved_output_paths=reserved_output_paths,
                )
                ext = os.path.splitext(file_path)[1].lower()

                if ext in (".doc", ".docx", ".xls", ".xlsx", ".xlsm", ".ppt", ".pptx") and os.path.getsize(file_path) == 0:
                    if not _same_file_path(file_path, output_path):
                        shutil.copy2(file_path, output_path)
                    count = 0
                elif ext == ".docx":
                    count = replace_in_docx(file_path, rules, output_path)
                elif ext in (".xlsx", ".xlsm"):
                    count = replace_in_xlsx(file_path, rules, output_path)
                elif ext == ".pptx":
                    count = replace_in_pptx(file_path, rules, output_path)
                elif ext in (".doc", ".xls", ".ppt"):
                    from legacy_office import (
                        LegacyOfficeSession,
                        replace_in_doc,
                        replace_in_ppt,
                        replace_in_xls,
                    )

                    if legacy_session is None:
                        candidate_session = LegacyOfficeSession()
                        try:
                            candidate_session.__enter__()
                        except Exception:
                            candidate_session.close()
                            raise
                        legacy_session = candidate_session
                    if ext == ".doc":
                        count = replace_in_doc(file_path, rules, output_path, legacy_session)
                    elif ext == ".xls":
                        count = replace_in_xls(file_path, rules, output_path, legacy_session)
                    else:
                        count = replace_in_ppt(file_path, rules, output_path, legacy_session)
                else:
                    errors.append(f"{filename}: 不支持的文件格式")
                    continue

                results[_result_key(results, filename, file_path)] = count
            except Exception as exc:
                errors.append(f"{filename}: {_format_processing_error(exc)}")
    finally:
        if legacy_session is not None:
            legacy_session.close()

    return results, "\n".join(errors) if errors else None
